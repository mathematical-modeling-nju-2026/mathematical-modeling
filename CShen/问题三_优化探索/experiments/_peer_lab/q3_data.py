"""Read the original attachments; all power samples refer to interval ends."""
from pathlib import Path
import hashlib
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
RAW = Path(r"d:\数学建模大赛\mathematical-modeling\C_yang\raw")
T, DT = 144, 1 / 6
HOURS = (0, 6, 12, 18)


def load_data():
    files = [RAW / "附件" / f"附件{i}.xlsx" for i in (1, 2, 3)]
    w = openpyxl.load_workbook(files[0], read_only=True, data_only=True)
    a = np.array([r[1:4] for r in w.active.iter_rows(min_row=2, values_only=True)], float)
    w.close()
    w = openpyxl.load_workbook(files[1], read_only=True, data_only=True)
    rows = [list(w[s].iter_rows(min_row=2, values_only=True))
            for s in ("小区负载", "光伏发电实际功率")]
    dates = [str(r[0])[:10] for r in rows[0]]
    assert [r[0] for r in rows[0]] == [r[0] for r in rows[1]]
    load, pv = (np.array([r[1:145] for r in rr], float) for rr in rows)
    w.close()
    w = openpyxl.load_workbook(files[2], read_only=True, data_only=True)
    rr = list(w.active.iter_rows(min_row=2, values_only=True))
    assert [str(r[1]) for r in rr[:4]] == ["0:00", "6:00", "12:00", "18:00"]
    external = np.array([r[2:26] for r in rr], float).reshape(365, 4, 24)
    w.close()
    assert load.shape == pv.shape == (365, T)
    assert all(np.isfinite(x).all() for x in (a, load, pv, external))
    assert min(load.min(), pv.min(), external.min()) >= 0
    return dict(price=a[:, 0], load=load, pv=pv, external=external,
                dates=dates, source_sha256={str(p.relative_to(RAW)): hashlib.sha256(p.read_bytes()).hexdigest()
                                          for p in files})


def label(t):
    return "24:00" if t == T else f"{t // 6:02d}:{t % 6 * 10:02d}"


def intervals(values, tol=1e-7):
    start = None
    for t in range(len(values) + 1):
        active = t < len(values) and values[t] > tol
        if active and start is None:
            start = t
        elif not active and start is not None:
            yield label(start) + "-" + label(t), float(values[start:t].sum())
            start = None
