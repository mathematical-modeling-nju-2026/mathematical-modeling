"""Read the original attachments; all power samples refer to interval ends."""
from pathlib import Path
import os
import hashlib
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
T, DT = 144, 1 / 6
HOURS = (0, 6, 12, 18)


def attachment_dir():
    """直接定位 C 题附件，避免依赖 C_yang/raw 的项目结构。"""
    configured = os.environ.get("CUMCM_C_ATTACHMENT_DIR")
    candidates = []
    if configured:
        candidates.append(Path(configured))
    candidates.extend([
        HERE / "附件",
        HERE / "data" / "附件",
        HERE.parents[1] / "CUMCM2026Problems" / "C题" / "附件",
    ])
    for candidate in candidates:
        if all((candidate / f"附件{i}.xlsx").is_file() for i in (1, 2, 3)):
            return candidate
    raise FileNotFoundError("找不到附件1.xlsx、附件2.xlsx、附件3.xlsx；请设置 CUMCM_C_ATTACHMENT_DIR。")


ATTACHMENTS = attachment_dir()


def load_data():
    files = [ATTACHMENTS / f"附件{i}.xlsx" for i in (1, 2, 3, 4)]
    w = openpyxl.load_workbook(files[0], read_only=True, data_only=True)
    a = np.array([r[1:4] for r in w.active.iter_rows(min_row=2, values_only=True)], float)
    w.close()
    w = openpyxl.load_workbook(files[1], read_only=True, data_only=True)
    rows = [list(w[s].iter_rows(min_row=2, values_only=True))
            for s in ("小区负载", "光伏发电实际功率")]
    dates = [str(r[0])[:10] for r in rows[0]]
    assert [r[0] for r in rows[0]] == [r[0] for r in rows[1]]
    load, pv = (np.array([r[1:145] for r in rr], float) for rr in rows)
    w.close()
    w = openpyxl.load_workbook(files[2], read_only=True, data_only=True)
    rr = list(w.active.iter_rows(min_row=2, values_only=True))
    assert [str(r[1]) for r in rr[:4]] == ["0:00", "6:00", "12:00", "18:00"]
    external = np.array([r[2:26] for r in rr], float).reshape(365, 4, 24)
    w.close()
    w = openpyxl.load_workbook(files[3], read_only=True, data_only=True)
    price_rt = np.array([r[1:145] for r in w.active.iter_rows(min_row=2, values_only=True)], float)
    w.close()
    assert load.shape == pv.shape == (365, T)
    assert price_rt.shape == (365, T)
    assert all(np.isfinite(x).all() for x in (a, load, pv, external, price_rt))
    assert min(load.min(), pv.min(), external.min(), price_rt.min()) >= 0
    return dict(price=a[:, 0], price_rt=price_rt, load=load, pv=pv, external=external,
                dates=dates, source_sha256={p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                                          for p in files})


def label(t):
    return "24:00" if t == T else f"{t // 6:02d}:{t % 6 * 10:02d}"


def intervals(values, tol=1e-7):
    start = None
    for t in range(len(values) + 1):
        active = t < len(values) and values[t] > tol
        if active and start is None:
            start = t
        elif not active and start is not None:
            yield label(start) + "-" + label(t), float(values[start:t].sum())
            start = None
