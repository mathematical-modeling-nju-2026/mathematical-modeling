"""按附件5 原模板重建全部 result*.xlsx（表头复用模板 + 环形映射填数）。

为什么要重建
    此前若干导出器把模板表头"纠正"成 00:00-00:10 … 23:50-24:00，
    与附件5 不一致，严格格式检查可能判为非模板。
    本脚本改为：**表头逐字复用模板原文**（含模板自身的 7:0-7:10 等写法），
    数值按官方参考实现的**环形映射**填入：
        out[i] = data[(i + 1) % 144]
    依据：CUMCM2026Problems/C题/问题一/make_result1.py

只改 Excel，不动任何计算结果（明细 CSV / JSON 均不触碰）。

用法
    python rebuild_result_workbooks.py --dry-run
    python rebuild_result_workbooks.py
    python rebuild_result_workbooks.py --only q1
"""
from __future__ import annotations

import argparse
import datetime as dt
import pathlib
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from template_layout import (block_labels, circular,  # noqa: E402
                             sheet_names, template_labels)

REPO = pathlib.Path(__file__).resolve().parents[2]
TPL = REPO / "data" / "附件" / "附件5"

# tag: (模板名, 明细相对路径, {工作表: (数值列, 费用列)}, 充电列, 放电列)
SPECS = {
    "q1": ("result1.xlsx", "schedule_detail.csv",
           {"计划购电量": ("g", "planned")}, "c", "d"),
    "q2": ("result2.xlsx", "schedule_detail.csv.gz",
           {"计划购电量": ("g_kwh", "planned_cost_yuan")}, "c_kwh", "d_kwh"),
    "q3": ("result3.xlsx", "schedule_detail.csv",
           {"计划购电量": ("original_kwh", "planned_cost_yuan"),
            "调整购电量": ("adjusted_kwh", "adjustment_net_yuan")},
           "charge_kwh", "discharge_kwh"),
    "q42": ("result4-2.xlsx", "variants/joint_price/schedule_detail.csv.gz",
            {"计划购电量": ("g_kwh", "planned_cost_yuan")}, "c_kwh", "d_kwh"),
    "q43": ("result4-3.xlsx", "schedule_detail.csv",
            {"计划购电量": ("original_kwh", "planned_cost_yuan"),
             "调整购电量": ("adjusted_kwh", "adjustment_net_yuan")},
            "charge_kwh", "discharge_kwh"),
}

# tag → (基线结果目录, def2 结果目录)
DIRS = {
    "q1":  ("question1/results", "question1/efficiency/results"),
    "q2":  ("question2/results", "question2/efficiency/results"),
    "q3":  ("question3/results", "question3/efficiency/results"),
    "q42": ("question4/part2/results", "question4/efficiency/part2/results"),
    "q43": ("question4/part3/results", "question4/efficiency/part3/results"),
}


def interval_label(k: int) -> str:
    """内部时段序号（0-based，区间末口径）→ 标签。"""
    m = k * 10
    return "0:00+1" if m >= 1440 else f"{m // 60}:{m % 60:02d}"


def merge_intervals(vals, tol=1e-7):
    """逐段紧急电量 → [(起, 止, 合计)]，标签为区间端点。"""
    out, t, n = [], 0, len(vals)
    while t < n:
        if vals[t] > tol:
            t0, q = t, 0.0
            while t < n and vals[t] > tol:
                q += float(vals[t]); t += 1
            out.append((interval_label(t0), interval_label(t), q))
        else:
            t += 1
    return out


def new_wb(tpl_name: str):
    names = sheet_names(TPL / tpl_name)
    wb = Workbook()
    wb.remove(wb.active)
    for n in names:
        wb.create_sheet(n)
    return wb


def style_book(wb) -> None:
    for ws in wb:
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (dt.datetime, dt.date)):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0000"


def build_q1(root: pathlib.Path, dry: bool) -> bool:
    """问题一：行式模板（145 行 × 2 列）。"""
    src = root / "schedule_detail.csv"
    if not src.exists():
        print(f"    跳过（缺 {src.name}）")
        return False
    d = pd.read_csv(src)
    cols = list(d.columns)
    tpl = TPL / "result1.xlsx"
    labels = template_labels(tpl, "计划购电量")
    wb = new_wb("result1.xlsx")

    g = d[cols[4]].to_numpy(float)          # 购电量(kWh)
    ws = wb["计划购电量"]
    ws.append(["时间段", "购电量"])
    for i, lab in enumerate(labels):
        ws.append([lab, round(float(g[(i + 1) % 144]), 6)])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16

    c, dd = d[cols[5]].to_numpy(float), d[cols[6]].to_numpy(float)
    E = d[cols[8]].to_numpy(float)
    ws2 = wb["充放电量"]
    ws2.append(["时间段", "充电量", "放电量", "时刻", "储电量"])
    blocks = block_labels(tpl, "充放电量")
    for b in range(6):
        a, z = b * 24, (b + 1) * 24
        ws2.append([blocks[b], round(float(c[a:z].sum()), 6),
                    round(float(dd[a:z].sum()), 6),
                    "0:00" if b == 0 else ("24:00" if b == 1 else None),
                    6000.0 if b == 0 else (float(E[-1]) if b == 1 else None)])
    for col, wd in zip("ABCDE", (16, 16, 16, 10, 16)):
        ws2.column_dimensions[col].width = wd

    style_book(wb)
    if not dry:
        wb.save(root / "result1.xlsx")
    print("    → result1.xlsx")
    return True


def build_wide(tag: str, root: pathlib.Path, dry: bool) -> bool:
    """问题 2/3/4：列式模板（日期 × 144 时段）。"""
    tpl_name, rel, sheets, ccol, dcol = SPECS[tag]
    src = root / rel
    if not src.exists():
        print(f"    跳过（缺 {rel}）")
        return False
    d = pd.read_csv(src)
    tpl = TPL / tpl_name
    labels = template_labels(tpl, "计划购电量")
    wb = new_wb(tpl_name)

    soc0_col = "energy_start_kwh" if "energy_start_kwh" in d.columns else "energy_before_kwh"
    soc1_col = "energy_end_kwh" if "energy_end_kwh" in d.columns else "energy_after_kwh"

    for sheet, (vcol, costcol) in sheets.items():
        ws = wb[sheet]
        ws.append(["日期\\时间"] + labels + ["全天购电量", "全天购电费"])
        for date, gdf in d.groupby("date", sort=False):
            vals = gdf[vcol].to_numpy(float)
            if len(vals) != 144:
                raise ValueError(f"{date} 段数 {len(vals)} != 144")
            ws.append([dt.datetime.fromisoformat(str(date))]
                      + [round(v, 6) for v in circular(vals)]
                      + [round(float(vals.sum()), 6),
                         round(float(gdf[costcol].sum()), 2)])
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 13

    ws2 = wb["充放电量"]
    ws2.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    blocks = block_labels(tpl, "充放电量")
    for date, gdf in d.groupby("date", sort=False):
        c = gdf[ccol].to_numpy(float)
        dd = gdf[dcol].to_numpy(float)
        e0 = float(gdf.iloc[0][soc0_col])
        e24 = float(gdf.iloc[-1][soc1_col])
        for b in range(6):
            a, z = b * 24, (b + 1) * 24
            ws2.append([dt.datetime.fromisoformat(str(date)) if b == 0 else None,
                        blocks[b], round(float(c[a:z].sum()), 6),
                        round(float(dd[a:z].sum()), 6),
                        "0:00" if b == 0 else ("24:00" if b == 1 else None),
                        e0 if b == 0 else (e24 if b == 1 else None)])
    for col, wd in zip("ABCDEF", (12, 14, 12, 12, 10, 12)):
        ws2.column_dimensions[col].width = wd

    ws3 = wb["紧急购电量"]
    ws3.append(["日期", "购电时间段", "购电量"])
    for date, gdf in d.groupby("date", sort=False):
        ivs = merge_intervals(gdf["emergency_kwh"].to_numpy(float))
        if not ivs:
            ws3.append([dt.datetime.fromisoformat(str(date)), "—", 0.0])
            continue
        for j, (a, z, q) in enumerate(ivs):
            ws3.append([dt.datetime.fromisoformat(str(date)) if j == 0 else None,
                        f"{a}-{z}", round(q, 6)])
    for col, wd in zip("ABC", (12, 22, 12)):
        ws3.column_dimensions[col].width = wd

    style_book(wb)
    if not dry:
        wb.save(root / f"result{tpl_name.split('result')[1]}")
    print(f"    → {tpl_name}")
    return True


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--only", default=None, choices=list(SPECS))
    args = ap.parse_args()
    dry = args.dry_run

    n = 0
    for tag in SPECS:
        if args.only and args.only != tag:
            continue
        for label, rel in (("基线", DIRS[tag][0]), ("def2", DIRS[tag][1])):
            root = REPO / rel
            print(f"== {tag} / {label}\n   {rel}")
            ok = build_q1(root, dry) if tag == "q1" else build_wide(tag, root, dry)
            n += int(ok)
    print()
    print(f"[dry-run] 待重建 {n} 个文件（未写入）" if dry else f"已重建 {n} 个 xlsx")


if __name__ == "__main__":
    main()
