"""Run comparable scenario-window experiments without changing the source B.

Default: four fixed-weight strategies plus the prespecified weekly selector.
--extended adds uniform 28/42-day windows and a six-candidate selector.
"""
from datetime import datetime, timedelta
from pathlib import Path
import argparse
import hashlib
import json
import sys
import time
import numpy as np
import pandas as pd
from rolling_window import HERE, SOURCE, Inputs, BY_NAME, PRIMARY_NAMES, CANDIDATES, prequential_scores, choose_week

# 附件5 模板的表头与「环形映射」规则（见 common/efficiency/template_layout.py）
_ETA_DIR = Path(__file__).resolve().parents[2] / "common" / "efficiency"
if str(_ETA_DIR) not in sys.path:
    sys.path.insert(0, str(_ETA_DIR))
from template_layout import (TPL2, block_labels, circular,  # noqa: E402
                             template_labels)


def label(slot):
    return '24:00' if slot == 144 else f'{slot//6:02d}:{slot%6*10:02d}'


def warmup(inputs):
    """Single common baseline January, preserving the original initial condition."""
    energy = 6000.
    rows = []
    for day in range(31):
        p, _, _ = inputs.solve_day(day, energy, BY_NAME['uniform56'])
        rows.append(dict(day=day,energy_start_kwh=energy,energy_end_kwh=float(p['E'][143])))
        energy = float(p['E'][143])
    return energy, rows


def simulate(inputs, name, initial, scores=None, names=PRIMARY_NAMES, first_residual=0, report_days=334):
    energy = initial
    rows, days, selections = [], [], []
    started = time.time()
    for day in range(31, min(365,31+report_days)):
        date = (datetime(2025,1,1)+timedelta(days=day)).strftime('%Y-%m-%d')
        if name.startswith('online'):
            selected, last_observed, means = choose_week(day, scores, names)
        else:
            selected, last_observed, means = name, day-1, []
        candidate = BY_NAME[selected]
        p, weights, indices = inputs.solve_day(day, energy, candidate, first_residual)
        g,c,d,E = (p[k][:144] for k in ('g','c','d','E'))
        before = np.r_[energy,E[:-1]]
        z = g+d-c
        emergency = np.maximum(inputs.actual[day]-z,0)
        dump = np.maximum(z-inputs.actual[day],0)
        planned_cost = inputs.price*g
        emergency_cost = 5*inputs.price*emergency
        total = planned_cost+emergency_cost
        for t in range(144):
            rows.append(dict(date=date,slot=t,start=label(t),end=label(t+1),price_yuan_per_kwh=inputs.price[t],
                             forecast_net_kwh=inputs.point[day,t],actual_net_kwh=inputs.actual[day,t],
                             g_kwh=g[t],c_kwh=c[t],d_kwh=d[t],energy_before_kwh=before[t],energy_after_kwh=E[t],
                             z_kwh=z[t],emergency_kwh=emergency[t],dump_kwh=dump[t],
                             planned_cost_yuan=planned_cost[t],emergency_cost_yuan=emergency_cost[t],total_cost_yuan=total[t]))
        days.append(dict(date=date,planned_cost_yuan=planned_cost.sum(),emergency_cost_yuan=emergency_cost.sum(),
                         total_cost_yuan=total.sum(),g_kwh=g.sum(),c_kwh=c.sum(),d_kwh=d.sum(),emergency_kwh=emergency.sum(),
                         dump_kwh=dump.sum(),energy_start_kwh=energy,energy_end_kwh=E[-1],chosen_candidate=selected,
                         sample_count=len(indices),effective_sample_count=1/np.sum(weights**2),
                         max_history_day=(datetime(2025,1,1)+timedelta(days=int(indices[-1]))).strftime('%Y-%m-%d') if len(indices) else '',
                         selection_last_observed_day=(datetime(2025,1,1)+timedelta(days=last_observed)).strftime('%Y-%m-%d')))
        selections.append(dict(date=date,candidate=selected,score_cutoff_day=last_observed,
                               candidate_names=list(names),mean_past_pinball_scores=means))
        energy = float(E[-1])
        if (day-31)%60==0 or day==min(364,30+report_days):
            print(f'{name:16s} {date} {sum(r["total_cost_yuan"] for r in days)/1e4:.3f} wan {time.time()-started:.1f}s',flush=True)
    detail,daily = pd.DataFrame(rows),pd.DataFrame(days)
    summary=dict(name=name,days=len(daily),initial_energy_kwh=initial,final_energy_kwh=energy,
                 first_residual_index=first_residual,seconds=time.time()-started)
    for col in ('planned_cost_yuan','emergency_cost_yuan','total_cost_yuan','g_kwh','c_kwh','d_kwh','emergency_kwh','dump_kwh'):
        summary[col]=float(daily[col].sum())
    summary['h1_cost_yuan']=float(daily.loc[daily.date<'2025-07-01','total_cost_yuan'].sum())
    summary['h2_cost_yuan']=float(daily.loc[daily.date>='2025-07-01','total_cost_yuan'].sum())
    return detail,daily,summary,selections


def export_workbook(detail, daily, path):
    import openpyxl
    from openpyxl.styles import Font
    # 模板表头（左端点，首格 0:10-0:20、末格 0:00+1-0:10+1）+ 环形映射：
    # 模板第 i 格 ← 内部时段 (i+1) % 144。依据官方问题一 make_result1.py。
    from template_layout import template_labels, circular, block_labels, TPL2
    tpl_labels = template_labels(TPL2, '计划购电量')
    blocks6 = block_labels(TPL2, '充放电量')
    w=openpyxl.Workbook(); ws=w.active; ws.title='计划购电量'
    ws.append(['日期\\时间']+tpl_labels+['全天购电量','全天购电费'])
    for date,g in detail.groupby('date',sort=False):
        vals = circular(g.g_kwh.to_numpy(float))
        ws.append([datetime.fromisoformat(date)]+[round(v,6) for v in vals]
                  +[g.g_kwh.sum(),g.planned_cost_yuan.sum()])
    ws.freeze_panes='B2'
    ws=w.create_sheet('充放电量'); ws.append(['日期','时间段','充电量','放电量','时刻','储电量'])
    for date,g in detail.groupby('date',sort=False):
        for b in range(6):
            block=g.iloc[b*24:(b+1)*24]
            ws.append([datetime.fromisoformat(date) if b==0 else None,blocks6[b],
                       block.c_kwh.sum(),block.d_kwh.sum(),'0:00' if b==0 else ('24:00' if b==1 else None),
                       g.iloc[0].energy_before_kwh if b==0 else (g.iloc[-1].energy_after_kwh if b==1 else None)])
    ws=w.create_sheet('紧急购电量'); ws.append(['日期','购电时间段','购电量'])
    for date,g in detail.groupby('date',sort=False):
        e=g.emergency_kwh.to_numpy(); blocks=[]; t=0
        while t<144:
            if e[t]<=1e-7: t+=1; continue
            start=t
            while t<144 and e[t]>1e-7: t+=1
            blocks.append((label(start)+'-'+label(t),e[start:t].sum()))
        for i,(span,amount) in enumerate(blocks or [('—',0.)]):
            ws.append([datetime.fromisoformat(date) if i==0 else None,span,amount])
    for ws in w:
        for cell in ws[1]: cell.font=Font(bold=True)
        ws.column_dimensions['A'].width=14
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,datetime):cell.number_format='yyyy-mm-dd'
                elif isinstance(cell.value,(int,float)):cell.number_format='0.0000'
    w.save(path)


def save_variant(out,name,result):
    detail,daily,summary,selections=result
    folder=out/'variants'/name;folder.mkdir(parents=True,exist_ok=True)
    detail.to_csv(folder/'schedule_detail.csv.gz',index=False,encoding='utf-8-sig',compression='gzip')
    daily.to_csv(folder/'daily_summary.csv',index=False,encoding='utf-8-sig')
    (folder/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    (folder/'selection_log.json').write_text(json.dumps(selections,ensure_ascii=False,indent=2),encoding='utf-8')
    export_workbook(detail,daily,folder/'result2.xlsx')
    return summary


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--extended',action='store_true')
    ap.add_argument('--first-residual',type=int,default=0)
    ap.add_argument('--initial',type=float,default=None)
    ap.add_argument('--days',type=int,default=334)
    ap.add_argument('--out',type=Path,default=HERE)
    ap.add_argument('--only',nargs='+',default=None)
    args=ap.parse_args()
    out=args.out if args.out.is_absolute() else HERE/args.out; out.mkdir(parents=True,exist_ok=True)
    before=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    inputs=Inputs()
    initial,january=warmup(inputs) if args.initial is None else (args.initial,[])
    names=tuple(c.name for c in CANDIDATES) if args.extended else PRIMARY_NAMES
    scores,quantiles=prequential_scores(inputs,names,args.first_residual)
    scoreframe=pd.DataFrame(scores,columns=names)
    scoreframe.insert(0,'date',pd.date_range('2025-01-01',periods=365))
    scoreframe.to_csv(out/'prequential_scores.csv',index=False,encoding='utf-8-sig')
    np.savez_compressed(out/'prequential_quantiles.npz',values=quantiles,candidates=np.array(names))
    (out/'common_january.json').write_text(json.dumps(january,indent=2),encoding='utf-8')
    modes=args.only or list(names)+['online_extended' if args.extended else 'online_primary']
    summaries=[]
    for name in modes:
        result=simulate(inputs,name,initial,scores,names,args.first_residual,args.days)
        summaries.append(save_variant(out,name,result))
    comparison=pd.DataFrame(summaries)
    if 'uniform56' in comparison.name.values:
        baseline=float(comparison.loc[comparison.name=='uniform56','total_cost_yuan'].iloc[0])
        comparison['saving_yuan']=baseline-comparison.total_cost_yuan
        comparison['saving_percent']=comparison.saving_yuan/baseline*100
    comparison.to_csv(out/'comparison.csv',index=False,encoding='utf-8-sig')
    assert before==hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    metadata=dict(source_file=str(SOURCE),source_sha256=before,first_residual=args.first_residual,
                  initial_soc=initial,year_end_terminal_soc=6000.,year_end_executed=args.days>=334,
                  horizon_days=2,candidates=[c.__dict__ for c in CANDIDATES if c.name in names],
                  online_rule='Every 7 days, select minimum mean 0.8-pinball loss over previous 28 completed days; min14 valid days; fixed-priority ties, uniform56 fallback.',
                  fixed_parameter_comparison='Exploratory retrospective comparison, not an untouched out-of-sample winner claim.',
                  inherited_forecaster='Original sw4/drift5 + t7/drift28 without changes, including early attachment1 fallback and negative PV predictions.',
                  python=sys.version)
    (out/'run_metadata.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    print(comparison.to_string(index=False),flush=True)


if __name__=='__main__': main()
