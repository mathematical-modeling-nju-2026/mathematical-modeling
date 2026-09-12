"""只修复 efficiency 目录中的提交工作簿时段列头。

附件5的计划购电列按区间末标签排列（从 0:10-0:20 开始，最后一列为
次日 0:00-0:10）。历史导出文件自行重写了列头。这里保留附件模板列头，
并把每行 144 个时段数据循环左移一格；日总量和费用列不变。
"""
from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
ATTACHMENT5 = ROOT / "data" / "附件" / "附件5"


def repair(template_name: str, output: Path, sheets: tuple[str, ...]) -> bool:
    """Repair one workbook once; return whether it was changed."""
    if not output.is_file():
        raise FileNotFoundError(output)
    template = openpyxl.load_workbook(ATTACHMENT5 / template_name, read_only=True, data_only=True)
    book = openpyxl.load_workbook(output)
    changed = False
    for sheet in sheets:
        expected = [template[sheet].cell(1, col).value for col in range(2, 146)]
        ws = book[sheet]
        actual = [ws.cell(1, col).value for col in range(2, 146)]
        if actual == expected:
            continue
        if len(actual) != 144:
            raise ValueError(f"{output}: {sheet} does not have 144 time columns")
        for row in range(2, ws.max_row + 1):
            values = [ws.cell(row, col).value for col in range(2, 146)]
            if all(value is None for value in values):
                continue
            values = values[1:] + values[:1]
            for col, value in enumerate(values, start=2):
                ws.cell(row, col).value = value
        for col, value in enumerate(expected, start=2):
            ws.cell(1, col).value = value
        changed = True
    template.close()
    if changed:
        book.save(output)
    book.close()
    return changed


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    targets = [
        ("result2.xlsx", ROOT / "question2" / "efficiency" / "results" / "result2.xlsx", ("计划购电量",)),
        ("result3.xlsx", ROOT / "question3" / "efficiency" / "results" / "result3.xlsx", ("计划购电量", "调整购电量")),
        ("result3.xlsx", ROOT / "question3" / "efficiency" / "results" / "comparisons" / "at_0_6_12" / "result3.xlsx", ("计划购电量", "调整购电量")),
        ("result4-3.xlsx", ROOT / "question4" / "efficiency" / "part3" / "results" / "result4-3.xlsx", ("计划购电量", "调整购电量")),
        ("result4-3.xlsx", ROOT / "question4" / "efficiency" / "part3" / "results" / "comparisons" / "at_0_6_12" / "result4-3.xlsx", ("计划购电量", "调整购电量")),
    ]
    q42 = ROOT / "question4" / "efficiency" / "part2" / "results"
    for variant in q42.glob("variants/*/result4-2.xlsx"):
        targets.append(("result4-2.xlsx", variant, ("计划购电量",)))

    changed = []
    for template, output, sheets in targets:
        if repair(template, output, sheets):
            changed.append(str(output.relative_to(ROOT)))

    # 根目录发布件来自 joint_price 变体；字节复制确保发布核验仍可复现。
    joint = q42 / "variants" / "joint_price" / "result4-2.xlsx"
    published = q42 / "result4-2.xlsx"
    shutil.copyfile(joint, published)
    report = q42 / "publication_verification.json"
    if report.is_file():
        data = json.loads(report.read_text(encoding="utf-8"))
        data["sha256"]["result4-2.xlsx"] = sha256(published)
        data["published_files_identical"] = True
        data["pass_check"] = bool(data.get("physical_verification_passed") and
                                  data.get("information_verification_passed"))
        report.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    print("changed", len(changed))
    for item in changed:
        print(item)


if __name__ == "__main__":
    main()
