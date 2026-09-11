"""问题2 数据层：读取附件1（电价/典型日）与附件2（全年实际负载与光伏）。

单位口径
    - 附件1/附件2 的负载与光伏均为功率 kW，时间分辨率 10 分钟
    - 每段电量(kWh) = 功率(kW) × Δt，Δt = 1/6 h
    - 附件1 的行标签是"区间末"（0:10 表示 00:00-00:10），共 144 段

问题2 的数据口径（依据题面）
    - 电价：只用附件1 的电价列，且"每天的电价相同" -> 全年复用同一条 144 段电价曲线
    - 负载/光伏：用附件2 的实际值（2025.1.1-12.31，365 天）
    - 净负载 N = 负载 - 光伏（kW 或 kWh，同比例）

附件5/result2.xlsx 的表头时间标签与附件1 的行标签存在一格错位
（模板首列为 '0:10-0:20'，末列为 '0:00+1-0:10+1'），二者是一一对应的，
本模块直接从模板读取标签，保证导出文件与模板逐列对齐。
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl

DT = 1.0 / 6.0          # 每段小时数
T = 144                 # 每天时段数
TAU = 10                # 每段分钟数

ROOT = Path(__file__).resolve().parents[1]
ATT1 = ROOT / "raw" / "附件" / "附件1.xlsx"
ATT2 = ROOT / "raw" / "附件" / "附件2.xlsx"
TPL2 = ROOT / "raw" / "附件" / "附件5" / "result2.xlsx"


def _series_from_row(row):
    """把附件2 的一行（日期 + 144 个数值）转成 float 数组。"""
    vals = np.array([float(v) for v in row[1:1 + T]], dtype=float)
    assert vals.size == T, f"期望 {T} 个时段，实际 {vals.size}"
    return vals


def load_attachment1():
    """附件1：返回 (labels, price, load_kW, pv_kW)。labels 为区间末标签。"""
    ws = openpyxl.load_workbook(ATT1, data_only=True).worksheets[0]
    labels, price, load, pv = [], [], [], []
    for t, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        labels.append(_label(t))
        price.append(float(row[1]))
        load.append(float(row[2]))
        pv.append(float(row[3]))
    assert len(price) == T, f"附件1 期望 {T} 段，实际 {len(price)} 段"
    return labels, np.array(price), np.array(load), np.array(pv)


def _label(t):
    """第 t 段（1-based）的区间末标签，t=144 -> '24:00'。"""
    m = t * TAU
    return f"{m // 60:02d}:{m % 60:02d}"


def load_attachment2():
    """附件2：返回 (dates, load_kW, pv_kW)，形状均为 (365, 144)。

    日期取自"小区负载"工作表第一列，已核对与"光伏发电实际功率"工作表一致。
    """
    wb = openpyxl.load_workbook(ATT2, data_only=True)
    ws_l = wb["小区负载"]
    ws_v = wb["光伏发电实际功率"]

    dates, load, pv = [], [], []
    for rl, rv in zip(ws_l.iter_rows(min_row=2, values_only=True),
                      ws_v.iter_rows(min_row=2, values_only=True)):
        if rl[0] is None:
            continue
        dates.append(rl[0])
        load.append(_series_from_row(rl))
        pv.append(_series_from_row(rv))
        assert rl[0] == rv[0], "两个工作表的日期不一致"

    load = np.vstack(load)
    pv = np.vstack(pv)
    assert load.shape == (365, T), f"附件2 形状异常: {load.shape}"
    return dates, load, pv


def template_labels():
    """从附件5/result2.xlsx 读取 144 个时间标签，保证导出与模板对齐。"""
    ws = openpyxl.load_workbook(TPL2, data_only=True)["计划购电量"]
    head = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    labs = [str(x) for x in head[1:1 + T]]
    assert len(labs) == T, f"模板标签数 {len(labs)}"
    return labs


def net_load(load_kW, pv_kW):
    """净负载（kW），等于负载减光伏。"""
    return load_kW - pv_kW


def to_energy(power_kW):
    """功率 -> 每段电量(kWh)。"""
    return np.asarray(power_kW, dtype=float) * DT


def weekday_index(dates):
    """返回每天的星期序号（周一=0）。"""
    return np.array([d.weekday() for d in dates], dtype=int)


if __name__ == "__main__":
    labs1, price, l1, v1 = load_attachment1()
    dates, L, V = load_attachment2()
    tl = template_labels()

    print("=" * 66)
    print("附件1（典型日）")
    print("=" * 66)
    print(f"  时段数   : {len(price)}")
    print(f"  电价     : {price.min():.4f} ~ {price.max():.4f} 元/kWh  均值 {price.mean():.4f}")
    print(f"  负载     : 全天 {l1.sum() * DT:,.2f} kWh  峰值 {l1.max():,.2f} kW")
    print(f"  光伏     : 全天 {v1.sum() * DT:,.2f} kWh  峰值 {v1.max():,.2f} kW")

    print("\n" + "=" * 66)
    print("附件2（2025 全年实际）")
    print("=" * 66)
    print(f"  天数     : {L.shape[0]}  ({dates[0].date()} ~ {dates[-1].date()})")
    print(f"  负载     : 全年 {L.sum() * DT / 1000:,.1f} MWh  日均 {L.sum() * DT / 365:,.1f} kWh")
    print(f"  光伏     : 全年 {V.sum() * DT / 1000:,.1f} MWh  日均 {V.sum() * DT / 365:,.1f} kWh")
    N = net_load(L, V)
    print(f"  净负载   : 日均 {N.sum() * DT / 365:,.1f} kWh  "
          f"最小 {N.min() * DT:.1f} kWh/段  最大 {N.max() * DT:.1f} kWh/段")

    print("\n附件1 是否为'典型日'（与附件2 日均比对）")
    print(f"  全天负载  附件1 {l1.sum() * DT:,.2f}  vs  附件2 日均 {L.sum() * DT / 365:,.2f} kWh"
          f"   偏差 {(l1.sum() - L.sum() / 365) / (L.sum() / 365) * 100:+.3f}%")
    print(f"  全天光伏  附件1 {v1.sum() * DT:,.2f}  vs  附件2 日均 {V.sum() * DT / 365:,.2f} kWh"
          f"   偏差 {(v1.sum() - V.sum() / 365) / (V.sum() / 365) * 100:+.3f}%")

    print("\n模板时间标签（前3/后3）:", tl[:3], "...", tl[-3:])
    print("附件1 行标签（前3/后3）:", labs1[:3], "...", labs1[-3:])
