"""第一问：仅从 ../raw 读取输入，解 LP 与 MILP，输出完整策略和论文表格。

时段采用右端点口径：附件的 00:10 表示 00:00--00:10 的平均功率。
变量块依次为 g, c, d, E, w；MILP 另加二元变量 z。
所有优化变量的单位都是 kWh，z 除外。
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime, time
from hashlib import sha256
import json
from pathlib import Path
import platform
import sys
from time import perf_counter

import numpy as np
import openpyxl
import pandas as pd
import scipy
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import hstack, lil_matrix, vstack

HERE = Path(__file__).resolve().parent
RAW = HERE.parent / "raw"
DATA = RAW / "附件" / "附件1.xlsx"
TEMPLATE = RAW / "附件" / "附件5" / "result1.xlsx"
TOL = 1e-6


@dataclass(frozen=True)
class Parameters:
    dt_hours: float = 1 / 6
    eta_charge: float = 0.9
    eta_discharge: float = 0.9
    capacity_kwh: float = 12000.0
    energy_min_kwh: float = 1200.0
    energy_max_kwh: float = 10800.0
    energy_initial_kwh: float = 6000.0
    power_max_kw: float = 5000.0


def hhmm(minutes: int) -> str:
    return f"{minutes // 60}:{minutes % 60:02d}"


def end_minute(value) -> int:
    """兼容 Excel 时间、数字日分数和 0:00+1 文本，不依赖行号猜测时间。"""
    if isinstance(value, (time, datetime)):
        if value.second or value.microsecond:
            raise ValueError(f"时间精度不是整分钟：{value}")
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)):
        m = float(value) * 1440
        if abs(m - round(m)) > 1e-5:
            raise ValueError(f"时间不是整分钟：{value}")
        return round(m)
    text = str(value).strip()
    day = 1440 if text.endswith("+1") else 0
    text = text.removesuffix("+1")
    parts = text.split(":")
    if len(parts) != 2:
        raise ValueError(f"无法解析时间：{value}")
    hour, minute = map(int, parts)
    if not 0 <= minute < 60 or not 0 <= hour <= 24:
        raise ValueError(f"非法时间：{value}")
    return 60 * hour + minute + day


def read_input():
    data = pd.read_excel(DATA, sheet_name=0)
    columns = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    if list(data.columns) != columns or len(data) != 144:
        raise ValueError(f"附件结构不符：{data.shape}, {list(data.columns)}")
    ends = np.array([end_minute(x) for x in data["时间"]])
    if not np.array_equal(ends, np.arange(10, 1441, 10)):
        raise ValueError("时间序列必须从 00:10 到 24:00，连续、唯一且间隔 10 分钟")
    numeric = data[columns[1:]].apply(pd.to_numeric, errors="raise").to_numpy(float)
    if not np.isfinite(numeric).all() or np.any(numeric < 0):
        raise ValueError("电价、负载、光伏含缺失、非有限值或负值")
    p, load, pv = numeric.T
    audit = {
        "row_count": len(data), "step_minutes": 10,
        "time_order_continuous_unique": True,
        "missing_numeric_count": int(np.isnan(numeric).sum()),
        "nonfinite_numeric_count": int((~np.isfinite(numeric)).sum()),
        "negative_numeric_count": int((numeric < 0).sum()),
        "price_range_yuan_per_kwh": [float(p.min()), float(p.max())],
        "load_range_kw": [float(load.min()), float(load.max())],
        "pv_range_kw": [float(pv.min()), float(pv.max())],
        "time_convention": "row timestamp is interval end; mean power over preceding 10 minutes",
        "source_sha256": {str(x.relative_to(RAW)): sha256(x.read_bytes()).hexdigest()
                          for x in (DATA, TEMPLATE, RAW / "C题.pdf")},
    }
    return ends, p, load, pv, audit


def build_model(price, load_kwh, pv_kwh, param):
    n = len(price)
    objective = np.zeros(5 * n)
    objective[:n] = price
    lower = np.zeros(5 * n)
    upper = np.full(5 * n, np.inf)
    upper[n:3*n] = param.power_max_kw * param.dt_hours
    lower[3*n:4*n] = param.energy_min_kwh
    upper[3*n:4*n] = param.energy_max_kwh
    upper[4*n:5*n] = pv_kwh
    # 两组逐时约束 + 日末状态约束；E_0 为给定参数，不是变量。
    eq = lil_matrix((2*n + 1, 5*n))
    rhs = np.zeros(2*n + 1)
    for t in range(n):
        # g + V + d = L + c + w。
        eq[t, t], eq[t, n+t], eq[t, 2*n+t], eq[t, 4*n+t] = 1, -1, 1, -1
        rhs[t] = load_kwh[t] - pv_kwh[t]
        # E_t - E_{t-1} - eta_c*c + d/eta_d = 0。
        eq[n+t, n+t] = -param.eta_charge
        eq[n+t, 2*n+t] = 1 / param.eta_discharge
        eq[n+t, 3*n+t] = 1
        if t:
            eq[n+t, 3*n+t-1] = -1
        else:
            rhs[n+t] = param.energy_initial_kwh
    eq[-1, 4*n-1] = 1
    rhs[-1] = param.energy_initial_kwh
    return objective, lower, upper, eq.tocsr(), rhs


def solve_models(price, load_kwh, pv_kwh, param):
    n = len(price)
    objective, lower, upper, eq, rhs = build_model(price, load_kwh, pv_kwh, param)
    started = perf_counter()
    lp = linprog(objective, A_eq=eq, b_eq=rhs, bounds=list(zip(lower, upper)),
                 method="highs", options={"primal_feasibility_tolerance": 1e-8,
                                           "dual_feasibility_tolerance": 1e-8})
    lp_seconds = perf_counter() - started
    if not lp.success:
        raise RuntimeError(f"LP 求解失败：{lp.message}")

    # 显式 MILP，不依赖对 LP 解事后裁剪。
    switch = lil_matrix((2*n, 6*n))
    limit = param.power_max_kw * param.dt_hours
    for t in range(n):
        switch[t, n+t], switch[t, 5*n+t] = 1, -limit
        switch[n+t, 2*n+t], switch[n+t, 5*n+t] = 1, limit
    all_constraints = vstack([hstack([eq, lil_matrix((2*n+1, n))]), switch]).tocsr()
    constraint_lower = np.r_[rhs, np.full(2*n, -np.inf)]
    constraint_upper = np.r_[rhs, np.zeros(n), np.full(n, limit)]
    started = perf_counter()
    mip = milp(np.r_[objective, np.zeros(n)], integrality=np.r_[np.zeros(5*n), np.ones(n)],
               bounds=Bounds(np.r_[lower, np.zeros(n)], np.r_[upper, np.ones(n)]),
               constraints=LinearConstraint(all_constraints, constraint_lower, constraint_upper),
               options={"mip_rel_gap": 1e-9, "time_limit": 120.0})
    mip_seconds = perf_counter() - started
    if not mip.success:
        raise RuntimeError(f"MILP 未证实最优：{mip.message}")
    simultaneous_lp = int(np.sum((lp.x[n:2*n] > TOL) & (lp.x[2*n:3*n] > TOL)))
    # LP 已互斥且与 MILP 同价时保留 LP 调度；否则输出显式 MILP 调度。
    selected = "LP" if simultaneous_lp == 0 and abs(lp.fun - mip.fun) < TOL else "MILP"
    x = (lp.x if selected == "LP" else mip.x[:5*n]).copy()

    # LP 对偶证据：有限边界对偶目标 + 一阶平稳残差。
    finite_upper = np.isfinite(upper)
    dual_objective = float(rhs @ lp.eqlin.marginals + lower @ lp.lower.marginals
                           + upper[finite_upper] @ lp.upper.marginals[finite_upper])
    stationarity = objective - eq.T @ lp.eqlin.marginals - lp.lower.marginals - lp.upper.marginals
    info = {
        "selected_schedule": selected,
        "lp": {"status": int(lp.status), "message": lp.message,
               "objective_yuan": float(lp.fun), "seconds": lp_seconds,
               "simultaneous_charge_discharge_periods": simultaneous_lp,
               "dual_objective_yuan": dual_objective,
               "primal_dual_gap_yuan": abs(float(lp.fun) - dual_objective),
               "stationarity_max_abs": float(np.max(np.abs(stationarity)))},
        "milp": {"status": int(mip.status), "message": mip.message,
                 "objective_yuan": float(mip.fun), "seconds": mip_seconds,
                 "dual_bound_yuan": float(mip.mip_dual_bound),
                 "relative_gap": float(mip.mip_gap), "node_count": int(mip.mip_node_count),
                 "binary_max_deviation": float(np.max(abs(mip.x[5*n:] - np.rint(mip.x[5*n:]))))},
        "lp_milp_objective_difference_yuan": abs(float(lp.fun - mip.fun)),
    }
    return x, info


def verify_schedule(x, load_kwh, pv_kwh, price, param):
    """按物理公式重新计算，不使用优化器约束残差。"""
    g, c, d, e, w = np.split(x, 5)
    previous = np.r_[param.energy_initial_kwh, e[:-1]]
    reconstructed = param.energy_initial_kwh + np.cumsum(param.eta_charge*c - d/param.eta_discharge)
    limit = param.power_max_kw * param.dt_hours
    checks = {
        "energy_balance_max_abs_kwh": float(np.max(abs(g + pv_kwh + d - load_kwh - c - w))),
        "state_recursion_max_abs_kwh": float(np.max(abs(e - previous - param.eta_charge*c + d/param.eta_discharge))),
        "state_cumulative_max_abs_kwh": float(np.max(abs(e - reconstructed))),
        "terminal_error_kwh": abs(float(e[-1] - param.energy_initial_kwh)),
        "bound_violation_max_kwh": float(max(0, -g.min(), -c.min(), -d.min(), -w.min(),
            (w-pv_kwh).max(), c.max()-limit, d.max()-limit,
            param.energy_min_kwh-e.min(), e.max()-param.energy_max_kwh)),
        "simultaneous_charge_discharge_periods": int(((c > TOL) & (d > TOL)).sum()),
        "cyclic_charge_discharge_error_kwh": abs(float(d.sum() - param.eta_charge*param.eta_discharge*c.sum())),
    }
    if not np.isfinite(x).all() or any(v > TOL for v in checks.values()):
        raise AssertionError(f"调度未通过物理核验：{checks}")
    checks["passed"] = True
    checks["tolerance_kwh"] = TOL
    return checks


def write_template(frame, blocks, param):
    book = openpyxl.load_workbook(TEMPLATE)
    if book.sheetnames != ["计划购电量", "充放电量"]:
        raise ValueError("结果模板工作表不符")
    ws = book["计划购电量"]
    old_labels = [ws.cell(i, 1).value for i in range(2, 146)]
    for row, record in enumerate(frame.itertuples(index=False), 2):
        ws.cell(row, 1, record.interval)
        ws.cell(row, 2, float(record.grid_kwh)).number_format = "0.000000"
    battery = book["充放电量"]
    for row, block in enumerate(blocks, 2):
        battery.cell(row, 2, block["charge_kwh"]).number_format = "0.000000"
        battery.cell(row, 3, block["discharge_kwh"]).number_format = "0.000000"
    battery["E2"] = param.energy_initial_kwh
    battery["E3"] = float(frame.energy_end_kwh.iloc[-1])
    book.save(HERE / "result1.xlsx")
    return {"original_first_label": old_labels[0], "original_last_label": old_labels[-1],
            "output_first_label": frame.interval.iloc[0], "output_last_label": frame.interval.iloc[-1],
            "reason": "align template labels with the selected interval-end convention and the 00:00--24:00 horizon"}


def write_paper_tables(frame, summary):
    total = summary["totals"]
    chosen = {r["interval"]: r["grid_kwh"] for r in summary["specified_grid"]}
    lines = ["# 第一问论文用表", "", "所有电量单位均为 kWh，购电费单位为元。展示保留两位小数；文件 result1.xlsx 和 dispatch_detail.csv 保留完整计算精度。", "",
             "**表 1　微网在指定时间段的购电量及全天的购电量和购电费**", "",
             "| 时间段 | 购电量 | 时间段 | 购电量 | 时间段 | 购电量 |",
             "|---|---:|---|---:|---|---:|"]
    for hours in ((10, 12, 14), (16, 18, 20)):
        cells = []
        for hour in hours:
            label = f"{hour}:00-{hour}:10"
            cells += [label, f"{chosen[label]:.2f}"]
        lines.append("| " + " | ".join(cells) + " |")
    lines += [f"| 全天购电量 | {total['grid_kwh']:.2f} | 全天购电费 | {total['cost_yuan']:.2f} | — | — |", "",
              "**表 2　储能设备在指定时间段的充放电量及 0:00 和 24:00 的储电量**", "",
              "| 时间段 | 充电量 | 放电量 | 时间段 | 充电量 | 放电量 |",
              "|---|---:|---:|---|---:|---:|"]
    blocks = summary["four_hour_blocks"]
    for i in range(0, 6, 2):
        cells = []
        for block in blocks[i:i+2]:
            cells += [block["interval"], f"{block['charge_kwh']:.2f}", f"{block['discharge_kwh']:.2f}"]
        lines.append("| " + " | ".join(cells) + " |")
    lines += [f"| 0:00 储电量 | {summary['parameters']['energy_initial_kwh']:.2f} | — | 24:00 储电量 | {frame.energy_end_kwh.iloc[-1]:.2f} | — |", "",
              "注意：一个四小时段内的充电量和放电量都可能为正，表示其不同十分钟子时段分别充电、放电，不代表同时充放电。", ""]
    (HERE / "paper_tables.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    param = Parameters()
    ends, price, load_kw, pv_kw, audit = read_input()
    load_kwh, pv_kwh = load_kw * param.dt_hours, pv_kw * param.dt_hours
    x, solver_info = solve_models(price, load_kwh, pv_kwh, param)
    checks = verify_schedule(x, load_kwh, pv_kwh, price, param)
    # 仅清除求解器的接近零数值噪声；再次核验，绝不四舍五入再派发。
    x[np.abs(x) < 1e-9] = 0
    verify_schedule(x, load_kwh, pv_kwh, price, param)
    g, c, d, e, w = np.split(x, 5)
    baseline = np.maximum(load_kwh-pv_kwh, 0)
    baseline_curtailment = np.maximum(pv_kwh-load_kwh, 0)
    intervals = [f"{hhmm(int(m)-10)}-{hhmm(int(m))}" for m in ends]
    frame = pd.DataFrame({
        "period": np.arange(1, 145), "interval": intervals,
        "start_minute": ends-10, "end_minute": ends, "price_yuan_per_kwh": price,
        "load_kw": load_kw, "pv_kw": pv_kw, "load_kwh": load_kwh, "pv_kwh": pv_kwh,
        "grid_kwh": g, "charge_kwh": c, "discharge_kwh": d, "curtailment_kwh": w,
        "energy_start_kwh": np.r_[param.energy_initial_kwh, e[:-1]], "energy_end_kwh": e,
        "soc_end": e / param.capacity_kwh, "cost_yuan": price*g,
        "no_storage_grid_kwh": baseline, "no_storage_cost_yuan": price*baseline,
        "mode": np.where(c > TOL, "充电", np.where(d > TOL, "放电", "待机")),
    })
    frame.to_csv(HERE / "dispatch_detail.csv", index=False, encoding="utf-8-sig")
    blocks = []
    for start in range(0, 24, 4):
        block = frame[(frame.start_minute >= start*60) & (frame.start_minute < (start+4)*60)]
        blocks.append({"interval": f"{start}:00-{start+4}:00",
                       **{col: float(block[col].sum()) for col in
                          ("grid_kwh", "charge_kwh", "discharge_kwh", "curtailment_kwh", "cost_yuan")},
                       "energy_end_kwh": float(block.energy_end_kwh.iloc[-1])})
    cost = float(price @ g)
    baseline_cost = float(price @ baseline)
    totals = {
        "cost_yuan": cost, "grid_kwh": float(g.sum()), "load_kwh": float(load_kwh.sum()),
        "pv_kwh": float(pv_kwh.sum()), "charge_kwh": float(c.sum()), "discharge_kwh": float(d.sum()),
        "curtailment_kwh": float(w.sum()), "storage_loss_kwh": float(c.sum()-d.sum()),
        "energy_min_kwh": float(e.min()), "energy_max_kwh": float(e.max()),
        "energy_final_kwh": float(e[-1]), "baseline_cost_yuan": baseline_cost,
        "baseline_grid_kwh": float(baseline.sum()), "baseline_curtailment_kwh": float(baseline_curtailment.sum()),
        "saving_yuan": baseline_cost-cost, "saving_percent": 100*(baseline_cost-cost)/baseline_cost,
        "pv_utilization_percent": 100*(1-w.sum()/pv_kwh.sum()),
        "baseline_pv_utilization_percent": 100*(1-baseline_curtailment.sum()/pv_kwh.sum()),
        "charge_periods": int((c > TOL).sum()), "discharge_periods": int((d > TOL).sum()),
        "idle_periods": int(((c <= TOL) & (d <= TOL)).sum()),
    }
    summary = {
        "parameters": asdict(param), "data_audit": audit, "solver": solver_info,
        "physical_verification": checks, "totals": totals, "four_hour_blocks": blocks,
        "specified_grid": frame[frame.start_minute.isin(np.array([10, 12, 14, 16, 18, 20])*60)]
                          [["interval", "grid_kwh"]].to_dict(orient="records"),
        "template_time_label_adjustment": write_template(frame, blocks, param),
        "environment": {"python": platform.python_version(), "numpy": np.__version__,
                        "pandas": pd.__version__, "scipy": scipy.__version__, "openpyxl": openpyxl.__version__},
    }
    (HERE / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_paper_tables(frame, summary)
    print(json.dumps({"totals": totals, "solver": solver_info, "checks": checks,
                      "four_hour_blocks": blocks, "specified_grid": summary["specified_grid"]},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
