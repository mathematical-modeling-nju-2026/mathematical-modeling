"""独立复核落盘结果；不导入 solve.py，不复用其约束矩阵。

重新读取原始 Excel 与结果文件，按能量守恒逐行复算；
另用消去 E、w 后的累计电量 LP 核验最优目标值。
"""
from hashlib import sha256
import json
from pathlib import Path
import sys

import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import linprog

HERE = Path(__file__).resolve().parent
RAW = HERE.parent / "raw"


def main():
    summary = json.loads((HERE / "summary.json").read_text(encoding="utf-8"))
    frame = pd.read_csv(HERE / "dispatch_detail.csv")
    param = summary["parameters"]
    source = openpyxl.load_workbook(RAW / "附件" / "附件1.xlsx", data_only=True)
    rows = list(source.active.iter_rows(min_row=2, values_only=True))
    price, load_kw, pv_kw = np.array([row[1:4] for row in rows], dtype=float).T
    dt, eta_c, eta_d = param["dt_hours"], param["eta_charge"], param["eta_discharge"]
    load, pv = load_kw*dt, pv_kw*dt
    initial, minimum, maximum = (param[k] for k in
                                 ("energy_initial_kwh", "energy_min_kwh", "energy_max_kwh"))
    limit = param["power_max_kw"]*dt
    n = len(rows)
    if len(frame) != n or n != 144:
        raise AssertionError("时段数错误")
    errors = {}

    def same(name, left, right, tolerance=1e-6):
        left, right = np.asarray(left, float), np.asarray(right, float)
        if not np.isfinite(left).all() or not np.isfinite(right).all():
            raise AssertionError(f"{name} 含非有限值")
        error = float(np.max(np.abs(left-right)))
        errors[name] = error
        if error > tolerance:
            raise AssertionError(f"{name}: {error} > {tolerance}")

    same("source_price", frame.price_yuan_per_kwh, price)
    same("source_load_kw", frame.load_kw, load_kw)
    same("source_pv_kw", frame.pv_kw, pv_kw)
    same("load_unit_conversion", frame.load_kwh, load)
    same("pv_unit_conversion", frame.pv_kwh, pv)
    same("time_starts", frame.start_minute, np.arange(0, 1440, 10))
    same("time_ends", frame.end_minute, np.arange(10, 1441, 10))
    g, c, d, w, e = (frame[col].to_numpy() for col in
                     ("grid_kwh", "charge_kwh", "discharge_kwh", "curtailment_kwh", "energy_end_kwh"))
    same("balance_kwh", g+pv+d, load+c+w)
    reconstructed = initial + np.cumsum(eta_c*c-d/eta_d)
    same("state_cumulative_kwh", e, reconstructed)
    same("state_starts_kwh", frame.energy_start_kwh, np.r_[initial, e[:-1]])
    same("terminal_kwh", e[-1], initial)
    same("soc", frame.soc_end, e/param["capacity_kwh"])
    same("period_cost_yuan", frame.cost_yuan, price*g)
    same("baseline_grid_kwh", frame.no_storage_grid_kwh, np.maximum(load-pv, 0))
    same("baseline_period_cost_yuan", frame.no_storage_cost_yuan, price*np.maximum(load-pv, 0))
    same("daily_cost_yuan", summary["totals"]["cost_yuan"], price@g)
    same("daily_grid_kwh", summary["totals"]["grid_kwh"], g.sum())
    same("daily_energy_conservation_kwh", g.sum()+pv.sum(), load.sum()+w.sum()+c.sum()-d.sum())
    violation = max(0, -g.min(), -c.min(), -d.min(), -w.min(), (w-pv).max(),
                    c.max()-limit, d.max()-limit, minimum-e.min(), e.max()-maximum)
    same("bound_violation_kwh", violation, 0)
    simultaneous = int(((c > 1e-6) & (d > 1e-6)).sum())
    same("simultaneous_periods", simultaneous, 0)

    book = openpyxl.load_workbook(HERE / "result1.xlsx", data_only=True)
    if book.sheetnames != ["计划购电量", "充放电量"]:
        raise AssertionError("模板工作表名称发生变化")
    grid_sheet, battery_sheet = book.worksheets
    if (grid_sheet.max_row, grid_sheet.max_column) != (145, 2):
        raise AssertionError("计划购电表形状错误")
    if (battery_sheet.max_row, battery_sheet.max_column) != (7, 5):
        raise AssertionError("充放电表形状错误")
    same("workbook_grid_kwh", [grid_sheet.cell(i+2, 2).value for i in range(n)], g)
    for i in range(n):
        start, end = 10*i, 10*(i+1)
        interval = f"{start//60}:{start%60:02d}-{end//60}:{end%60:02d}"
        if grid_sheet.cell(i+2, 1).value != interval or frame.interval.iloc[i] != interval:
            raise AssertionError(f"第 {i+1} 个时段标签不一致")
    for j in range(6):
        section = slice(24*j, 24*(j+1))
        same(f"workbook_block_{j+1}_charge_kwh", battery_sheet.cell(j+2, 2).value, c[section].sum())
        same(f"workbook_block_{j+1}_discharge_kwh", battery_sheet.cell(j+2, 3).value, d[section].sum())
        block = summary["four_hour_blocks"][j]
        same(f"summary_block_{j+1}_charge_kwh", block["charge_kwh"], c[section].sum())
        same(f"summary_block_{j+1}_discharge_kwh", block["discharge_kwh"], d[section].sum())
    same("workbook_initial_kwh", battery_sheet["E2"].value, initial)
    same("workbook_final_kwh", battery_sheet["E3"].value, e[-1])
    for item in summary["specified_grid"]:
        same(f"specified_{item['interval']}", item["grid_kwh"],
             frame.loc[frame.interval == item["interval"], "grid_kwh"].iloc[0])

    # 另一种建模：变量只保留 g,c,d；电池状态用充放电累计和表达。
    eye = np.eye(n)
    zeros = np.zeros((n, n))
    cumulative = np.tril(np.ones((n, n)))
    flow = np.hstack([eye, -eye, eye])
    state = np.hstack([zeros, eta_c*cumulative, -cumulative/eta_d])
    # 0 <= w = g+V+d-L-c <= V，即 L-V <= g-c+d <= L。
    a_ub = np.vstack([-flow, flow, state, -state])
    b_ub = np.r_[pv-load, load, np.full(n, maximum-initial), np.full(n, initial-minimum)]
    objective = np.r_[price, np.zeros(2*n)]
    result = linprog(objective, A_ub=a_ub, b_ub=b_ub,
                     A_eq=state[-1:], b_eq=[0.0],
                     bounds=[(0, None)]*n + [(0, limit)]*(2*n), method="highs")
    if not result.success:
        raise AssertionError(f"累计电量 LP 失败：{result.message}")
    same("independent_lp_cost_difference_yuan", result.fun, price@g)

    hashes_ok = all(sha256((RAW / name).read_bytes()).hexdigest() == digest
                    for name, digest in summary["data_audit"]["source_sha256"].items())
    if not hashes_ok:
        raise AssertionError("原始输入文件与求解时哈希值不同")
    report = {
        "passed": True, "tolerance": 1e-6, "check_count": len(errors),
        "max_abs_errors": errors, "source_hashes_unchanged": hashes_ok,
        "independent_formulation": "cumulative-state LP with g,c,d only; no solve.py import",
        "independent_lp_status": int(result.status), "independent_lp_message": result.message,
        "independent_lp_cost_yuan": float(result.fun),
        "conclusion": "Exported schedule is feasible with charge/discharge exclusivity and attains the independent LP lower bound within tolerance.",
    }
    (HERE / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
