"""Causal day-ahead forecasting and paired whole-day residual scenarios.

The two forecasting rules are selected using January only.  A prediction at
origin d uses observations strictly before d, even for the second forecast day.
All public forecast/scenario arrays use kWh per ten-minute interval.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

DT_HOURS = 1.0 / 6.0
SLOTS = 144
WARMUP_DAYS = 7
JANUARY_DAYS = 31
SCENARIO_LOOKBACK = 28
CANDIDATES = (
    "last_day", "mean7", "mean14", "weighted7", "weekday_last",
    "weekday_mean4", "trend7", "trend14", "weekday_trend",
)


@dataclass(frozen=True)
class YearData:
    dates: tuple[date, ...]
    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray
    audit: dict


@dataclass(frozen=True)
class ForecastBundle:
    load_kwh: np.ndarray  # origin, lead_day, slot; first seven origins are NaN
    pv_kwh: np.ndarray
    selected_methods: dict[str, str]
    validation: dict


@dataclass(frozen=True)
class ScenarioBundle:
    load_kwh: np.ndarray  # scenario, lead_day, slot
    pv_kwh: np.ndarray
    weights: np.ndarray
    history_indices: np.ndarray  # historical origins; last target < today's origin


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _minute(value) -> int:
    if isinstance(value, time):
        return 60 * value.hour + value.minute
    if str(value).strip() in {"0:00+1", "00:00+1", "24:00"}:
        return 1440
    if isinstance(value, str) and value.count(":") == 1:
        hour, minute = value.strip().split(":")
        if hour.isdigit() and minute.isdigit() and 0 <= int(hour) < 24 and 0 <= int(minute) < 60:
            return 60 * int(hour) + int(minute)
    raise ValueError(f"Unexpected interval endpoint: {value!r}")


def _describe(x: np.ndarray, dates: tuple[date, ...]) -> dict:
    return {
        "shape": list(x.shape), "missing_or_nonfinite_count": int((~np.isfinite(x)).sum()),
        "negative_count": int((x < 0).sum()), "zero_count": int((x == 0).sum()),
        "duplicate_entire_days": int(len(x) - len(np.unique(x, axis=0))),
        "minimum_kw": float(x.min()), "maximum_kw": float(x.max()),
        "mean_kw": float(x.mean()),
        "monthly_mean_kw": {
            str(m): float(x[[d.month == m for d in dates]].mean()) for m in range(1, 13)
        },
    }


def load_data(raw_dir: Path | str | None = None) -> YearData:
    """Read only raw/附件/附件1.xlsx and 附件2.xlsx, validating all time cells."""
    raw = Path(raw_dir) if raw_dir is not None else Path(__file__).resolve().parent.parent / "raw"
    a1, a2 = raw / "附件" / "附件1.xlsx", raw / "附件" / "附件2.xlsx"
    wb = load_workbook(a2, read_only=True, data_only=True)
    required = ("小区负载", "光伏发电实际功率")
    if tuple(wb.sheetnames) != required:
        raise ValueError(f"Unexpected attachment 2 sheets: {wb.sheetnames}")
    arrays, date_axes = [], []
    expected_minutes = np.arange(1, SLOTS + 1) * 10
    for name in required:
        rows = list(wb[name].values)
        if len(rows) != 366 or len(rows[0]) != 145:
            raise ValueError(f"Unexpected sheet shape for {name}")
        minutes = np.asarray([_minute(v) for v in rows[0][1:]])
        if not np.array_equal(minutes, expected_minutes):
            raise ValueError(f"Ten-minute endpoint sequence is invalid: {name}")
        date_axis = tuple(r[0].date() if isinstance(r[0], datetime) else r[0] for r in rows[1:])
        arr = np.asarray([r[1:] for r in rows[1:]], dtype=float)
        if not np.isfinite(arr).all() or np.any(arr < 0):
            raise ValueError(f"Missing, nonfinite or negative observations in {name}")
        arrays.append(arr)
        date_axes.append(date_axis)
    wb.close()
    dates = date_axes[0]
    expected_dates = tuple(date(2025, 1, 1) + timedelta(days=i) for i in range(365))
    if dates != expected_dates or date_axes[1] != dates:
        raise ValueError("Date axes must match every day of 2025 in increasing order")
    wb1 = load_workbook(a1, read_only=True, data_only=True)
    rows1 = list(wb1.worksheets[0].values)
    wb1.close()
    if len(rows1) != 145:
        raise ValueError("Attachment 1 must contain 144 time intervals")
    if not np.array_equal([_minute(r[0]) for r in rows1[1:]], expected_minutes):
        raise ValueError("Attachment 1 and attachment 2 interval endpoints differ")
    price = np.asarray([r[1] for r in rows1[1:]], dtype=float)
    if not np.isfinite(price).all() or np.any(price <= 0):
        raise ValueError("Electricity prices must be finite and positive")
    audit = {
        "inputs": {str(a1.relative_to(raw)): _sha256(a1), str(a2.relative_to(raw)): _sha256(a2)},
        "sheet_names": list(required), "dates": [dates[0].isoformat(), dates[-1].isoformat()],
        "number_of_days": len(dates), "slots_per_day": SLOTS,
        "time_convention": "Each power value is a ten-minute mean; 00:10 labels 00:00-00:10 and 0:00+1 labels 23:50-24:00.",
        "dt_hours": DT_HOURS, "load": _describe(arrays[0], dates),
        "pv": _describe(arrays[1], dates),
        "price": {"source": "attachment 1, repeated every day", "minimum_yuan_per_kwh": float(price.min()), "maximum_yuan_per_kwh": float(price.max())},
        "data_modifications": "None; only kW to kWh conversion for optimization.",
        "calendar_note": "Do not impose Saturday/Sunday low load: January observations exhibit a seven-day cycle with lower demand on Friday/Saturday.",
        "use_of_annual_summaries": "Descriptive audit only; future monthly statistics are never forecasting inputs.",
    }
    return YearData(dates, price, arrays[0], arrays[1], audit)


def _predict_series(history: np.ndarray, horizon_days: int, method: str) -> np.ndarray:
    """No future array is accepted: history contains only fully observed days."""
    n_days = len(history)
    if n_days < WARMUP_DAYS or horizon_days not in (1, 2):
        raise ValueError("At least seven historical days and a one/two-day horizon are required")
    out = []
    for h in range(horizon_days):
        if method == "last_day":
            pred = history[-1]
        elif method in ("mean7", "mean14"):
            pred = history[-int(method[4:]):].mean(axis=0)
        elif method == "weighted7":
            values = history[-7:]
            pred = np.average(values, axis=0, weights=np.arange(1, len(values) + 1))
        elif method == "weekday_last":
            pred = history[n_days + h - 7]
        elif method == "weekday_mean4":
            indices = np.arange(n_days + h - 7, max(-1, n_days + h - 29), -7)
            pred = history[indices].mean(axis=0)
        elif method in ("trend7", "trend14"):
            n = min(int(method[5:]), n_days)
            axis = np.arange(n, dtype=float)
            centered = axis - axis.mean()
            values = history[-n:]
            slope = centered @ values / (centered @ centered)
            pred = values.mean(axis=0) + (n + h - axis.mean()) * slope
        elif method == "weekday_trend":
            correction = (history[-7:].mean(axis=0) - history[-14:-7].mean(axis=0)) if n_days >= 14 else 0.0
            pred = history[n_days + h - 7] + correction
        else:
            raise ValueError(f"Unknown prediction method: {method}")
        out.append(np.maximum(pred, 0.0))
    return np.asarray(out)


def _metrics(errors: np.ndarray) -> dict:
    return {
        "mae_kw": float(np.mean(np.abs(errors))),
        "rmse_kw": float(np.sqrt(np.mean(errors ** 2))),
        "bias_prediction_minus_actual_kw": float(np.mean(errors)),
        "number_of_interval_predictions": int(errors.size),
    }


def select_methods(data: YearData) -> tuple[dict[str, str], dict]:
    """Prequential comparison within January; February-December never consulted."""
    selected, scores = {}, {}
    for label, values in (("load", data.load_kw[:31]), ("pv", data.pv_kw[:31])):
        scores[label] = {}
        for method in CANDIDATES:
            errors = []
            for origin in range(14, JANUARY_DAYS):
                horizon = min(2, JANUARY_DAYS - origin)
                pred = _predict_series(values[:origin], horizon, method)
                errors.append((pred - values[origin:origin + horizon]).ravel())
            scores[label][method] = _metrics(np.concatenate(errors))
        selected[label] = min(CANDIDATES, key=lambda m: scores[label][m]["mae_kw"])
    return selected, scores


def forecast_at(data: YearData, day_index: int, horizon_days: int = 2,
                methods: dict[str, str] | None = None) -> tuple[np.ndarray, np.ndarray]:
    """Return (load, pv) forecasts, each [H,144] kWh, using data[:day_index]."""
    if not WARMUP_DAYS <= day_index <= len(data.dates):
        raise ValueError("day_index outside available historical range")
    if methods is None:
        if day_index < JANUARY_DAYS:
            raise ValueError("Pass explicit rules for January hindcasts; automatic January selection is not yet available")
        methods, _ = select_methods(data)
    return (
        _predict_series(data.load_kw[:day_index], horizon_days, methods["load"]) * DT_HOURS,
        _predict_series(data.pv_kw[:day_index], horizon_days, methods["pv"]) * DT_HOURS,
    )


def build_forecasts(data: YearData) -> ForecastBundle:
    methods, january_scores = select_methods(data)
    load = np.full((len(data.dates), 2, SLOTS), np.nan)
    pv = np.full_like(load, np.nan)
    for origin in range(WARMUP_DAYS, len(data.dates)):
        load[origin], pv[origin] = forecast_at(data, origin, 2, methods)
    validation = {
        "selection_period": "2025-01-15 through 2025-01-31; two leads included only if the target date is still January",
        "selection_metric": "Pooled one- and two-day-ahead MAE in kW; tie broken by fixed candidate order",
        "candidates": list(CANDIDATES), "january_comparison": january_scores,
        "selected_methods": methods,
        "method_descriptions": {
            "weekday_mean4": "Arithmetic mean at each slot of the latest up to four observed days sharing the target's weekday (7,14,21,28-day offsets).",
            "trend14": "Separate OLS line for each slot using the latest 14 fully observed days; extrapolate one and two days ahead, clip at zero.",
        },
        "test_period": "2025-02-01 through 2025-12-31",
        "test_rule": "Forecast forms and lookback constants frozen after January; historical observations update each day.",
        "january_residual_caveat": "Initial residuals are fixed-rule hindcasts using only each origin's past features; the choice among these rules used January validation, so initial residual calibration is not an independent test set.",
        "scenario_rule": "Most recent up to 28 complete paired residual blocks, starting no earlier than January 8; uniform weights; load and PV residuals stay paired across all 144 slots and both forecast leads; negative scenario powers clipped at zero.",
        "test_metrics": {}, "test_monthly_metrics_one_day_ahead": {},
    }
    for label, pred, actual in (("load", load, data.load_kw), ("pv", pv, data.pv_kw),
                                ("net_load", load - pv, data.load_kw - data.pv_kw)):
        validation["test_metrics"][label] = {}
        for h in (0, 1):
            origins = np.arange(JANUARY_DAYS, len(data.dates) - h)
            err = pred[origins, h] / DT_HOURS - actual[origins + h]
            validation["test_metrics"][label][f"lead_{h + 1}_day"] = _metrics(err)
        validation["test_monthly_metrics_one_day_ahead"][label] = {
            str(m): _metrics(pred[[i for i, d in enumerate(data.dates) if d.month == m], 0] / DT_HOURS
                             - actual[[i for i, d in enumerate(data.dates) if d.month == m]])
            for m in range(2, 13)
        }
    return ForecastBundle(load, pv, methods, validation)


def make_scenarios(data: YearData, bundle: ForecastBundle, day_index: int,
                   horizon_days: int = 2, lookback: int = SCENARIO_LOOKBACK) -> ScenarioBundle:
    """Shift the current forecast by fully observed historical joint residuals."""
    if horizon_days not in (1, 2) or lookback < 1:
        raise ValueError("Expected a one/two-day horizon and positive lookback")
    last = day_index - horizon_days
    origins = np.arange(max(WARMUP_DAYS, last - lookback + 1), last + 1, dtype=int)
    if not len(origins) or day_index >= len(data.dates):
        raise ValueError("No complete historical residual blocks for this origin")
    load_scenarios, pv_scenarios = [], []
    for j in origins:
        load_residual = data.load_kw[j:j + horizon_days] * DT_HOURS - bundle.load_kwh[j, :horizon_days]
        pv_residual = data.pv_kw[j:j + horizon_days] * DT_HOURS - bundle.pv_kwh[j, :horizon_days]
        load_scenarios.append(np.maximum(bundle.load_kwh[day_index, :horizon_days] + load_residual, 0.0))
        pv_scenarios.append(np.maximum(bundle.pv_kwh[day_index, :horizon_days] + pv_residual, 0.0))
    return ScenarioBundle(np.asarray(load_scenarios), np.asarray(pv_scenarios),
                          np.full(len(origins), 1.0 / len(origins)), origins)


def check_causal_invariance(data: YearData, bundle: ForecastBundle) -> dict:
    """Replace current/future observations: forecasts and scenario sets must not move."""
    records = []
    for origin in (31, 180, 364):
        altered_load, altered_pv = data.load_kw.copy(), data.pv_kw.copy()
        altered_load[origin:] = 1_234_567.0
        altered_pv[origin:] = 7_654_321.0
        altered = YearData(data.dates, data.price, altered_load, altered_pv, data.audit)
        horizon = min(2, len(data.dates) - origin)
        pred = forecast_at(altered, origin, horizon, bundle.selected_methods)
        forecast_change = max(float(np.max(np.abs(pred[0] - bundle.load_kwh[origin, :horizon]))),
                              float(np.max(np.abs(pred[1] - bundle.pv_kwh[origin, :horizon]))))
        before = make_scenarios(data, bundle, origin, horizon)
        after = make_scenarios(altered, bundle, origin, horizon)
        scenario_change = max(float(np.max(np.abs(before.load_kwh - after.load_kwh))),
                              float(np.max(np.abs(before.pv_kwh - after.pv_kwh))))
        assert forecast_change == scenario_change == 0.0
        assert np.max(before.history_indices + horizon - 1) < origin
        records.append({"origin_date": data.dates[origin].isoformat(), "horizon_days": horizon,
                        "forecast_max_change_kwh": forecast_change, "scenario_max_change_kwh": scenario_change,
                        "number_of_scenarios": len(before.weights),
                        "last_observed_residual_date": data.dates[int(before.history_indices[-1] + horizon - 1)].isoformat()})
    altered_load, altered_pv = data.load_kw.copy(), data.pv_kw.copy()
    altered_load[31:], altered_pv[31:] = 1_234_567.0, 7_654_321.0
    changed = YearData(data.dates, data.price, altered_load, altered_pv, data.audit)
    changed_methods, changed_scores = select_methods(changed)
    original_methods, original_scores = select_methods(data)
    assert changed_methods == original_methods and changed_scores == original_scores
    return {"future_perturbation_checks": records, "method_selection_unchanged_after_replacing_all_test_data": True}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=Path(__file__).resolve().parent.parent / "raw")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    data = load_data(args.raw_dir)
    bundle = build_forecasts(data)
    bundle.validation["causality_checks"] = check_causal_invariance(data, bundle)
    (args.output_dir / "data_audit.json").write_text(json.dumps(data.audit, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output_dir / "forecast_validation.json").write_text(json.dumps(bundle.validation, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"selected_methods": bundle.selected_methods, "test_metrics": bundle.validation["test_metrics"],
                      "causality_checks": bundle.validation["causality_checks"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
