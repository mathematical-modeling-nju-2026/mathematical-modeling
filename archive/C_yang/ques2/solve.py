"""Run Q2 chronological backtest and fill raw attachment 5/result2.xlsx.

Business inputs: raw attachments 1, 2 and result2 template only.
January is the calibration period with idle storage. The selected policy is
stochastic_48h, chosen before observing February-December costs.
"""
from copy import copy
from dataclasses import asdict
from datetime import datetime, time, timedelta
from hashlib import sha256
import json
from pathlib import Path
import platform
import sys
from time import perf_counter

import numpy as np
import openpyxl
import pandas as pd
import scipy

from forecast import load_data, build_forecasts, make_scenarios, check_causal_invariance
from optimization import PARAM, solve_dispatch, execute_day, empirical_quantile

HERE = Path(__file__).resolve().parent
RAW = HERE.parent / "raw"
TEMPLATE = RAW / "附件" / "附件5" / "result2.xlsx"
METHODS = ("stochastic_48h", "deterministic_48h", "no_storage", "stochastic_24h")
SELECTED_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
TOL = 1e-6


def time_label(minutes):
    return f"{int(minutes)//60}:{int(minutes)%60:02d}"


INTERVALS = [f"{time_label(10*i)}-{time_label(10*(i+1))}" for i in range(144)]


def emergency_segments(frame):
    records = []
    for date, day in frame.groupby("date", sort=True):
        q = day.emergency_kwh.to_numpy()
        active = q > TOL
        changes = np.diff(np.r_[False, active, False].astype(int))
        starts, ends = np.where(changes == 1)[0], np.where(changes == -1)[0]
        if not len(starts):
            records.append({"date": date, "interval": "无", "emergency_kwh": 0.0,
                            "start_period": 0, "end_period": 0})
        for start, end in zip(starts, ends):
            records.append({"date": date, "interval": f"{time_label(start*10)}-{time_label(end*10)}",
                            "emergency_kwh": float(q[start:end].sum()),
                            "start_period": int(start+1), "end_period": int(end)})
    return pd.DataFrame(records)


def write_workbook(detail, daily, segments):
    wb = openpyxl.load_workbook(TEMPLATE)
    if wb.sheetnames != ["计划购电量", "充放电量", "紧急购电量"]:
        raise ValueError("Unexpected result2 workbook sheets")
    grid, battery, urgent = wb.worksheets
    old_labels = [grid.cell(1, j+2).value for j in range(144)]
    for j, interval in enumerate(INTERVALS, 2):
        grid.cell(1, j, interval)
    daily_selected = daily[daily.method == METHODS[0]].set_index("date")
    for row, (date, day) in enumerate(detail.groupby("date", sort=True), 2):
        grid.cell(row, 1, datetime.fromisoformat(date)).number_format = "yyyy-mm-dd"
        for j, value in enumerate(day.grid_plan_kwh, 2):
            grid.cell(row, j, float(value)).number_format = "0.000000"
        grid.cell(row, 146, float(day.grid_plan_kwh.sum())).number_format = "0.000000"
        # The worksheet is the day-ahead purchase ledger: emergency charges are
        # separately itemized in daily_summary.csv and the written report.
        grid.cell(row, 147, float(day.planned_cost_yuan.sum())).number_format = "0.000000"
    grid.freeze_panes = "B2"
    styles = [[copy(battery.cell(i, j)._style) for j in range(1, 7)] for i in range(2, 8)]
    battery.delete_rows(2, battery.max_row-1)
    for day_index, (date, day) in enumerate(detail.groupby("date", sort=True)):
        for block in range(6):
            row = 2+6*day_index+block
            for col in range(1, 7):
                battery.cell(row, col)._style = copy(styles[block][col-1])
            section = day.iloc[block*24:(block+1)*24]
            if block == 0:
                battery.cell(row, 1, datetime.fromisoformat(date)).number_format = "yyyy-mm-dd"
            battery.cell(row, 2, f"{4*block}:00-{4*(block+1)}:00")
            battery.cell(row, 3, float(section.charge_kwh.sum())).number_format = "0.000000"
            battery.cell(row, 4, float(section.discharge_kwh.sum())).number_format = "0.000000"
            if block == 0:
                battery.cell(row, 5, time(0, 0)).number_format = "h:mm"
                battery.cell(row, 6, float(day.energy_start_kwh.iloc[0])).number_format = "0.000000"
            elif block == 1:
                battery.cell(row, 5, "24:00")
                battery.cell(row, 6, float(day.energy_end_kwh.iloc[-1])).number_format = "0.000000"
    battery.freeze_panes = "C2"
    urgent_styles = [copy(urgent.cell(2, col)._style) for col in range(1, 4)]
    urgent.delete_rows(2, urgent.max_row-1)
    previous_date = None
    for row, item in enumerate(segments.itertuples(index=False), 2):
        for col in range(1, 4):
            urgent.cell(row, col)._style = copy(urgent_styles[col-1])
        if item.date != previous_date:
            urgent.cell(row, 1, datetime.fromisoformat(item.date)).number_format = "yyyy-mm-dd"
        urgent.cell(row, 2, item.interval)
        urgent.cell(row, 3, item.emergency_kwh).number_format = "0.000000"
        previous_date = item.date
    urgent.freeze_panes = "B2"
    urgent.column_dimensions["B"].width = 25
    wb.save(HERE / "result2.xlsx")
    return {"original_first_interval": old_labels[0], "original_last_interval": old_labels[-1],
            "output_first_interval": INTERVALS[0], "output_last_interval": INTERVALS[-1],
            "grid_sheet_daily_cost_definition": "planned purchases only: sum(price * grid_plan_kwh); emergency and total actual costs are separate in daily_summary.csv",
            "emergency_aggregation": "merge adjacent within-day ten-minute intervals with emergency_kwh > 1e-6; no events => 无,0",
            "all_dates_expanded": True, "original_template_sha256": sha256(TEMPLATE.read_bytes()).hexdigest()}


def write_tables(detail, daily, segments, bounds):
    lines = ["# 第二问论文用表", "",
             "所有电量单位为 kWh，费用单位为元。表 1 的购电量与购电费指日前计划，和 result2.xlsx 的计划购电量工作表一致；紧急购电费及实际总费另列，避免混合口径。展示两位小数，底层文件保存完整精度。", ""]
    for date in SELECTED_DATES:
        day = detail[detail.date == date]
        row = daily[(daily.date == date) & (daily.method == METHODS[0])].iloc[0]
        lines += [f"**{date}：表 1　指定时间段与全天计划购电**", "",
                  "| 时间段 | 购电量 | 时间段 | 购电量 | 时间段 | 购电量 |",
                  "|---|---:|---|---:|---|---:|"]
        for hours in ((10, 12, 14), (16, 18, 20)):
            cells = []
            for hour in hours:
                cells += [f"{hour}:00-{hour}:10", f"{day.iloc[hour*6].grid_plan_kwh:.2f}"]
            lines.append("| "+" | ".join(cells)+" |")
        lines += [f"| 全天计划购电量 | {row.grid_plan_kwh:.2f} | 全天计划购电费 | {row.planned_cost_yuan:.2f} | — | — |", "",
                  f"当天紧急购电 {row.emergency_kwh:.2f} kWh，紧急购电费 {row.emergency_cost_yuan:.2f} 元；**实际总购电费 {row.total_cost_yuan:.2f} 元**，计划与紧急电量合计 {row.grid_plan_kwh+row.emergency_kwh:.2f} kWh。", "",
                  f"**{date}：表 2　储能充放电及日初、日末储电量**", "",
                  "| 时间段 | 充电量 | 放电量 | 时间段 | 充电量 | 放电量 |",
                  "|---|---:|---:|---|---:|---:|"]
        for b in range(0, 6, 2):
            cells = []
            for j in (b, b+1):
                block = day.iloc[j*24:(j+1)*24]
                cells += [f"{j*4}:00-{(j+1)*4}:00", f"{block.charge_kwh.sum():.2f}", f"{block.discharge_kwh.sum():.2f}"]
            lines.append("| "+" | ".join(cells)+" |")
        lines += [f"| 0:00 储电量 | {row.energy_start_kwh:.2f} | — | 24:00 储电量 | {row.energy_end_kwh:.2f} | — |", ""]
    selected_segments = [segments[segments.date == date].reset_index(drop=True) for date in SELECTED_DATES]
    lines += ["**表 3　四个指定日期的紧急购电区间**", "",
              "| 3月20日时间段 | 购电量 | 6月21日时间段 | 购电量 | 9月23日时间段 | 购电量 | 12月21日时间段 | 购电量 |",
              "|---|---:|---|---:|---|---:|---|---:|"]
    for i in range(max(map(len, selected_segments))):
        cells = []
        for group in selected_segments:
            cells += [group.iloc[i].interval, f"{group.iloc[i].emergency_kwh:.2f}"] if i < len(group) else ["—", "—"]
        lines.append("| "+" | ".join(cells)+" |")
    lines += ["", "紧急购电连续区间已合并，全部日期记录另见 emergency_intervals.csv。四小时内充放电总量都为正，表示在不同十分钟时段发生。", "",
              "**四个指定日期的条件完全信息下界（仅用于评价）**", "",
              "| 日期 | 本策略实际费用 | 已知当天实际值的条件下界 | 费用差额 |",
              "|---|---:|---:|---:|"]
    for bound in bounds:
        lines.append(f"| {bound['date']} | {bound['actual_policy_cost_yuan']:.2f} | {bound['perfect_information_cost_yuan']:.2f} | {bound['gap_yuan']:.2f} |")
    lines += ["", "下界模型预先知道当天真实负载和光伏，并固定为本策略相同的日初、日末电量；不属于可以在当天 0:00 执行的方案，也不构成全年下界。", ""]
    (HERE / "paper_tables.md").write_text("\n".join(lines), encoding="utf-8")


def main():
    start = perf_counter()
    data = load_data(RAW)
    forecasts = build_forecasts(data)
    forecasts.validation["causality_checks"] = check_causal_invariance(data, forecasts)
    (HERE / "data_audit.json").write_text(json.dumps(data.audit, ensure_ascii=False, indent=2), encoding="utf-8")
    (HERE / "forecast_validation.json").write_text(json.dumps(forecasts.validation, ensure_ascii=False, indent=2), encoding="utf-8")
    states = {method: PARAM.initial_energy_kwh for method in METHODS}
    archive = {method: [] for method in METHODS}
    details, daily_rows, logs, bounds = [], [], [], []
    for day_index in range(31, len(data.dates)):
        date = data.dates[day_index].isoformat()
        h = min(2, len(data.dates)-day_index)
        scenarios = make_scenarios(data, forecasts, day_index, h, PARAM.scenario_lookback)
        s = len(scenarios.weights)
        net = (scenarios.load_kwh-scenarios.pv_kwh).reshape(s, -1)
        predicted_net = (forecasts.load_kwh[day_index, :h]-forecasts.pv_kwh[day_index, :h]).reshape(1, -1)
        actual_load = data.load_kw[day_index]*PARAM.dt_hours
        actual_pv = data.pv_kw[day_index]*PARAM.dt_hours
        prices = np.tile(data.price, h)
        for method in METHODS:
            initial = states[method]
            if method == "no_storage":
                quantile = empirical_quantile(net[:, :144], scenarios.weights)
                plan = {"grid": np.maximum(quantile, 0), "charge": np.zeros(144),
                        "discharge": np.zeros(144), "energy": np.full(144, initial),
                        "net_quantile": quantile}
            else:
                if method == "stochastic_48h":
                    plan = solve_dispatch(net, scenarios.weights, prices, initial)
                elif method == "deterministic_48h":
                    plan = solve_dispatch(predicted_net, np.ones(1), prices, initial)
                else:
                    plan = solve_dispatch(net[:, :144], scenarios.weights, data.price, initial)
                logs.append({"date": date, "method": method, **plan["log"],
                             "latest_input_date": (data.dates[day_index]-timedelta(days=1)).isoformat(),
                             "scenario_first_origin": data.dates[int(scenarios.history_indices[0])].isoformat(),
                             "scenario_last_origin": data.dates[int(scenarios.history_indices[-1])].isoformat(),
                             "scenario_last_observed_date": data.dates[int(scenarios.history_indices[-1]+h-1)].isoformat()})
            execution = execute_day(plan, actual_load, actual_pv, data.price, initial)
            states[method] = float(execution["energy_end_kwh"][-1])
            archive[method].append(execution)
            aggregate_fields = ("grid_plan_kwh", "charge_kwh", "discharge_kwh", "emergency_kwh",
                                "surplus_kwh", "planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan")
            daily_rows.append({"date": date, "method": method,
                               **{key: float(execution[key].sum()) for key in aggregate_fields},
                               "energy_start_kwh": initial, "energy_end_kwh": states[method],
                               "emergency_periods": int((execution["emergency_kwh"] > TOL).sum()),
                               "energy_min_kwh": float(execution["energy_end_kwh"].min()),
                               "energy_max_kwh": float(execution["energy_end_kwh"].max())})
            if method == METHODS[0]:
                detail = pd.DataFrame({"date": date, "period": np.arange(1, 145), "interval": INTERVALS,
                                       "price_yuan_per_kwh": data.price, "load_kwh": actual_load, "pv_kwh": actual_pv,
                                       "forecast_load_kwh": forecasts.load_kwh[day_index, 0],
                                       "forecast_pv_kwh": forecasts.pv_kwh[day_index, 0],
                                       "net_q80_kwh": empirical_quantile(net[:, :144], scenarios.weights),
                                       **execution})
                details.append(detail)
                if date in SELECTED_DATES:
                    oracle = solve_dispatch((actual_load-actual_pv)[None, :], np.ones(1), data.price,
                                            initial, states[method])
                    oracle_cost = float(data.price@oracle["grid"])
                    actual_cost = float(execution["total_cost_yuan"].sum())
                    bounds.append({"date": date, "energy_initial_kwh": initial,
                                   "energy_final_kwh": states[method], "actual_policy_cost_yuan": actual_cost,
                                   "perfect_information_cost_yuan": oracle_cost,
                                   "gap_yuan": actual_cost-oracle_cost,
                                   "label": "conditional single-day perfect-information lower bound, not deployable"})
        if day_index == 31 or day_index % 30 == 0 or day_index == 364:
            print(f"Backtest {date}: selected day cost {daily_rows[-4]['total_cost_yuan']:.2f} yuan, elapsed {perf_counter()-start:.1f}s", flush=True)
    detail = pd.concat(details, ignore_index=True)
    daily = pd.DataFrame(daily_rows)
    optimizer = pd.DataFrame(logs)
    segments = emergency_segments(detail)
    detail.to_csv(HERE / "dispatch_detail.csv", index=False, encoding="utf-8-sig")
    daily.to_csv(HERE / "daily_summary.csv", index=False, encoding="utf-8-sig")
    optimizer.to_csv(HERE / "optimizer_log.csv", index=False, encoding="utf-8-sig")
    segments.to_csv(HERE / "emergency_intervals.csv", index=False, encoding="utf-8-sig")
    fields = archive[METHODS[0]][0].keys()
    np.savez_compressed(HERE / "all_method_dispatch.npz", methods=np.array(METHODS),
                        dates=np.array([d.isoformat() for d in data.dates[31:]]),
                        **{key: np.array([[record[key] for record in archive[method]] for method in METHODS]) for key in fields})
    method_totals = {}
    for method in METHODS:
        group = daily[daily.method == method]
        method_totals[method] = {
            **{key: float(group[key].sum()) for key in aggregate_fields},
            "emergency_days": int((group.emergency_periods > 0).sum()),
            "emergency_periods": int(group.emergency_periods.sum()),
            "emergency_period_fraction": float(group.emergency_periods.sum()/(334*144)),
            "energy_initial_kwh": float(group.energy_start_kwh.iloc[0]),
            "energy_final_kwh": float(group.energy_end_kwh.iloc[-1]),
        }
    main_cost = method_totals[METHODS[0]]["total_cost_yuan"]
    comparisons = {method: {"saving_yuan": method_totals[method]["total_cost_yuan"]-main_cost,
                            "saving_percent": 100*(method_totals[method]["total_cost_yuan"]-main_cost)/method_totals[method]["total_cost_yuan"]}
                   for method in METHODS[1:]}
    summary = {
        "selected_method": METHODS[0], "method_selected_before_test_costs": True,
        "evaluation_dates": [data.dates[31].isoformat(), data.dates[-1].isoformat()],
        "evaluation_days": 334, "evaluation_periods": 334*144,
        "parameters": asdict(PARAM), "forecast_methods": forecasts.selected_methods,
        "january_state_policy": "Battery idle January 1-31: E=6000 continuously. January is calibration; costs before February are not scored.",
        "terminal_policy": "Daily rolling 48h reference terminal E=6000, execute first 24h, carry actual E to next day. Dec31 horizon=24h and final E=6000.",
        "benchmark_definitions": {
            "stochastic_48h": "selected fixed day-ahead g,c,d; finite-scenario expected planned+emergency fee minimization",
            "deterministic_48h": "same point forecasts and horizon, ignore uncertainty at planning time; settle true emergency costs",
            "no_storage": "same historical scenarios; empirical inverse-CDF 80% quantile purchases, battery idle",
            "stochastic_24h": "terminal/horizon sensitivity: same first-day scenario marginals, daily E returns to 6000"},
        "method_totals": method_totals, "comparisons_to_selected": comparisons,
        "specified_days": daily[(daily.method == METHODS[0]) & daily.date.isin(SELECTED_DATES)].to_dict(orient="records"),
        "conditional_perfect_information_bounds": bounds,
        "net_q80_observed_coverage": float(((detail.load_kwh-detail.pv_kwh) <= detail.net_q80_kwh+TOL).mean()),
        "selected_emergency_events": int((segments.interval != "无").sum()),
        "optimizer_audit": {"lp_solves": len(optimizer), "all_status_zero": bool((optimizer.lp_status == 0).all()),
                            "max_absolute_cleanup_cost_difference_yuan": float(optimizer.cleanup_objective_difference_yuan.abs().max()),
                            "max_cleanup_state_difference_kwh": float(optimizer.cleanup_state_max_abs_kwh.max()),
                            "simultaneous_before_cleanup_total": int(optimizer.simultaneous_periods_before_cleanup.sum()),
                            "simultaneous_after_cleanup_total": int(optimizer.simultaneous_periods_after_cleanup.sum())},
        "template": write_workbook(detail, daily, segments),
        "source_sha256": {**data.audit["inputs"], str(TEMPLATE.relative_to(RAW)): sha256(TEMPLATE.read_bytes()).hexdigest()},
        "environment": {"python": platform.python_version(), "numpy": np.__version__, "pandas": pd.__version__,
                        "scipy": scipy.__version__, "openpyxl": openpyxl.__version__},
        "runtime_seconds": perf_counter()-start,
    }
    (HERE / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_tables(detail, daily, segments, bounds)
    print(json.dumps({"method_totals": method_totals, "comparisons": comparisons,
                      "specified_days": summary["specified_days"], "optimizer": summary["optimizer_audit"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
