"""Independently audit exported question-2 results against the raw workbooks.

Energy, costs, emergency intervals, and workbook checks do not import the LP
model or reuse its constraint matrix. A separate causal perturbation audit may
import the forecasting / planning public API, solely to challenge information
access; it is not a second proof of global stochastic-policy optimality.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd

BASE = Path(__file__).resolve().parent
RAW = BASE.parent / "raw" / "附件"
TOL = 2e-5
ACTIVE_TOL = 1e-6
ETA = 0.9
LIMIT = 5000 / 6
CHECKS: list[dict[str, Any]] = []


def record(name: str, passed: bool, **evidence: Any) -> None:
    CHECKS.append({"check": name, "passed": bool(passed), **evidence})


def close(name: str, actual: Any, expected: Any, tol: float = TOL) -> None:
    a, b = np.asarray(actual, dtype=float), np.asarray(expected, dtype=float)
    if a.shape != b.shape:
        record(name, False, actual_shape=list(a.shape), expected_shape=list(b.shape))
        return
    finite = bool(np.isfinite(a).all() and np.isfinite(b).all())
    error = float(np.max(np.abs(a - b))) if a.size and finite else (0.0 if not a.size else None)
    record(name, finite and error is not None and error <= tol,
           max_absolute_error=error, tolerance=tol, number_of_values=int(a.size))


def hhmm(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def interval(start: int, end: int) -> str:
    return f"{hhmm(start * 10)}-{hhmm(end * 10)}"


def normalize_interval(value: Any) -> str:
    """Accept harmless zero-padding / dash differences in displayed times."""
    value = str(value).strip().replace("—", "-").replace("–", "-").replace(" ", "")
    parts = value.split("-")
    if len(parts) != 2:
        return value
    answer = []
    for part in parts:
        if part in {"0:00+1", "00:00+1"}:
            answer.append("24:00")
            continue
        try:
            hours, minutes = map(int, part.split(":"))
            answer.append(f"{hours:02d}:{minutes:02d}")
        except ValueError:
            return value
    return "-".join(answer)


def iso_date(value: Any) -> str:
    return pd.Timestamp(value).date().isoformat()


def raw_arrays() -> tuple[list[str], np.ndarray, np.ndarray, np.ndarray]:
    wb = openpyxl.load_workbook(RAW / "附件2.xlsx", read_only=True, data_only=True)
    values = [list(ws.iter_rows(min_row=2, values_only=True)) for ws in wb]
    record("raw_actual_workbook_two_sheets", len(values) == 2)
    dates = [iso_date(row[0]) for row in values[0]]
    record("raw_load_pv_date_alignment", dates == [iso_date(row[0]) for row in values[1]])
    load, pv = [np.asarray([row[1:145] for row in sheet], dtype=float) / 6 for sheet in values]
    record("raw_actual_dimensions", load.shape == pv.shape == (365, 144), shape=list(load.shape))
    record("raw_actual_finite_nonnegative", bool(np.isfinite(load).all() and np.isfinite(pv).all()
                                                   and (load >= 0).all() and (pv >= 0).all()))
    for ws in wb:
        headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))[1:]
        expected = [dt.time((k * 10) // 60, (k * 10) % 60) for k in range(1, 144)] + ["0:00+1"]
        record(f"raw_right_endpoint_times_{ws.title}", headers == expected)
    wb.close()
    wp = openpyxl.load_workbook(RAW / "附件1.xlsx", read_only=True, data_only=True)
    rows = list(wp.active.iter_rows(min_row=2, values_only=True))
    prices = np.asarray([row[1] for row in rows], dtype=float)
    record("raw_price_dimensions_and_positive", prices.shape == (144,) and bool((prices > 0).all()))
    wp.close()
    return dates, load, pv, prices


def check_trace(frame: pd.DataFrame, raw: tuple) -> dict[str, float]:
    dates, load, pv, prices = raw
    expected_dates = pd.date_range("2025-02-01", "2025-12-31").strftime("%Y-%m-%d").tolist()
    record("trace_334_days_144_periods", len(frame) == 334 * 144,
           actual_rows=len(frame), expected_rows=334 * 144)
    record("trace_date_order", frame.date.tolist() == np.repeat(expected_dates, 144).tolist())
    record("trace_period_order", frame.period.tolist() == np.tile(np.arange(1, 145), 334).tolist())
    expected_intervals = [interval(k, k + 1) for k in range(144)] * 334
    record("trace_right_endpoint_interval_labels",
           frame.interval.map(normalize_interval).tolist() == expected_intervals)
    numeric = frame.select_dtypes(include="number").to_numpy()
    record("trace_all_numeric_values_finite", bool(np.isfinite(numeric).all()))
    indices = [dates.index(date) for date in expected_dates]
    close("load_matches_raw_power_divided_by_6", frame.load_kwh, load[indices].ravel())
    close("pv_matches_raw_power_divided_by_6", frame.pv_kwh, pv[indices].ravel())
    close("price_matches_raw_daily_tariff", frame.price_yuan_per_kwh, np.tile(prices, 334))

    g = frame.grid_plan_kwh.to_numpy()
    c = frame.charge_kwh.to_numpy()
    d = frame.discharge_kwh.to_numpy()
    q = frame.emergency_kwh.to_numpy()
    w = frame.surplus_kwh.to_numpy()
    p = frame.price_yuan_per_kwh.to_numpy()
    soc_start = frame.energy_start_kwh.to_numpy()
    soc_end = frame.energy_end_kwh.to_numpy()
    for name, array in [("grid_plan", g), ("charge", c), ("discharge", d), ("emergency", q), ("surplus", w)]:
        record(f"{name}_nonnegative", bool((array >= -TOL).all()), minimum=float(array.min()))
    record("charge_external_power_limit", bool((c <= LIMIT + TOL).all()), maximum_kwh=float(c.max()), limit_kwh=LIMIT)
    record("discharge_external_power_limit", bool((d <= LIMIT + TOL).all()), maximum_kwh=float(d.max()), limit_kwh=LIMIT)
    simultaneous = (c > ACTIVE_TOL) & (d > ACTIVE_TOL)
    record("no_simultaneous_charge_discharge", not bool(simultaneous.any()), simultaneous_periods=int(simultaneous.sum()))
    reconstructed = 6000 + np.cumsum(ETA * c - d / ETA)
    close("soc_independent_full_year_cumulative_reconstruction", soc_end, reconstructed)
    close("soc_period_state_transition", soc_end, soc_start + ETA * c - d / ETA)
    close("soc_every_period_and_day_continuity", soc_start[1:], soc_end[:-1])
    close("february_1_start_after_january_standby", soc_start[:1], [6000])
    close("december_31_terminal_soc", soc_end[-1:], [6000])
    record("soc_1200_to_10800", bool((reconstructed >= 1200 - TOL).all() and (reconstructed <= 10800 + TOL).all()),
           minimum_kwh=float(reconstructed.min()), maximum_kwh=float(reconstructed.max()))
    gap = frame.load_kwh.to_numpy() + c - frame.pv_kwh.to_numpy() - d - g
    close("emergency_is_actual_positive_energy_gap", q, np.maximum(gap, 0))
    close("surplus_is_actual_negative_energy_gap", w, np.maximum(-gap, 0))
    close("actual_energy_conservation", g + frame.pv_kwh.to_numpy() + d + q,
          frame.load_kwh.to_numpy() + c + w)
    close("planned_cost_paid_for_entire_plan", frame.planned_cost_yuan, p * g)
    close("emergency_cost_at_5x_tariff", frame.emergency_cost_yuan, 5 * p * q)
    close("total_cost_is_plan_plus_emergency", frame.total_cost_yuan, p * g + 5 * p * q)
    return {
        "grid_plan_kwh": float(g.sum()), "emergency_kwh": float(q.sum()),
        "surplus_kwh": float(w.sum()), "charge_kwh": float(c.sum()), "discharge_kwh": float(d.sum()),
        "planned_cost_yuan": float(np.dot(p, g)), "emergency_cost_yuan": float(np.dot(5 * p, q)),
        "total_cost_yuan": float(np.dot(p, g) + np.dot(5 * p, q)),
        "minimum_soc_kwh": float(reconstructed.min()), "maximum_soc_kwh": float(reconstructed.max()),
    }


def check_daily(frame: pd.DataFrame) -> None:
    daily = pd.read_csv(BASE / "daily_summary.csv", dtype={"date": str})
    methods = {"stochastic_48h", "deterministic_48h", "no_storage", "stochastic_24h"}
    record("daily_comparison_methods", set(daily.method) == methods, methods=sorted(set(daily.method)))
    expected_dates = pd.date_range("2025-02-01", "2025-12-31").strftime("%Y-%m-%d").tolist()
    for method, data in daily.groupby("method", sort=False):
        record(f"daily_dates_{method}", data.date.tolist() == expected_dates)
        if {"planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan"}.issubset(data.columns):
            close(f"daily_cost_identity_{method}", data.total_cost_yuan,
                  data.planned_cost_yuan.to_numpy() + data.emergency_cost_yuan.to_numpy())
    selected = daily[daily.method == "stochastic_48h"].set_index("date")
    sums = frame.groupby("date", sort=False).sum(numeric_only=True)
    for col in ["grid_plan_kwh", "charge_kwh", "discharge_kwh", "emergency_kwh", "surplus_kwh",
                "planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan"]:
        if col in selected:
            close(f"daily_selected_{col}", selected.loc[sums.index, col], sums[col])
    for col, aggregation in [("energy_start_kwh", "first"), ("energy_end_kwh", "last")]:
        if col in selected:
            expected = getattr(frame.groupby("date", sort=False)[col], aggregation)()
            close(f"daily_selected_{col}", selected.loc[expected.index, col], expected)


def check_all_method_traces(frame: pd.DataFrame, raw: tuple) -> None:
    _, raw_load, raw_pv, prices = raw
    bundle = np.load(BASE / "all_method_dispatch.npz", allow_pickle=False)
    methods = bundle["methods"].tolist()
    dates = bundle["dates"].tolist()
    record("all_method_archive_method_order", methods == ["stochastic_48h", "deterministic_48h", "no_storage", "stochastic_24h"])
    record("all_method_archive_date_order", dates == pd.date_range("2025-02-01", "2025-12-31").strftime("%Y-%m-%d").tolist())
    fields = ["grid_plan_kwh", "charge_kwh", "discharge_kwh", "energy_start_kwh", "energy_end_kwh",
              "emergency_kwh", "surplus_kwh", "planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan"]
    record("all_method_archive_shapes", all(bundle[col].shape == (4, 334, 144) for col in fields))
    daily = pd.read_csv(BASE / "daily_summary.csv", dtype={"date": str})
    net = (raw_load[31:] - raw_pv[31:]).ravel()
    p = np.tile(prices, 334)
    for index, method in enumerate(methods):
        data = {col: bundle[col][index].ravel() for col in fields}
        for col in fields:
            if method == "stochastic_48h":
                close(f"archive_selected_csv_{col}", data[col], frame[col])
        record(f"{method}_all_trace_finite", all(np.isfinite(array).all() for array in data.values()))
        g, c, d, q, w = (data[col] for col in ["grid_plan_kwh", "charge_kwh", "discharge_kwh", "emergency_kwh", "surplus_kwh"])
        record(f"{method}_nonnegative_and_power_limits", bool(min(g.min(), c.min(), d.min(), q.min(), w.min()) >= -TOL
               and c.max() <= LIMIT + TOL and d.max() <= LIMIT + TOL))
        record(f"{method}_no_simultaneous_charge_discharge", not bool(((c > ACTIVE_TOL) & (d > ACTIVE_TOL)).any()))
        soc = 6000 + np.cumsum(ETA * c - d / ETA)
        close(f"{method}_soc_cumulative_reconstruction", data["energy_end_kwh"], soc)
        close(f"{method}_soc_all_start_states", data["energy_start_kwh"], np.r_[6000, soc[:-1]])
        record(f"{method}_soc_capacity_bounds", bool(soc.min() >= 1200 - TOL and soc.max() <= 10800 + TOL),
               minimum_kwh=float(soc.min()), maximum_kwh=float(soc.max()))
        close(f"{method}_terminal_6000", soc[-1:], [6000])
        close(f"{method}_emergency_from_raw_actual_gap", q, np.maximum(net + c - d - g, 0))
        close(f"{method}_surplus_from_raw_actual_gap", w, np.maximum(-net - c + d + g, 0))
        close(f"{method}_planned_cost_paid_for_entire_plan", data["planned_cost_yuan"], p * g)
        close(f"{method}_emergency_cost_5x", data["emergency_cost_yuan"], 5 * p * q)
        close(f"{method}_total_cost_identity", data["total_cost_yuan"], p * g + 5 * p * q)
        rows = daily[daily.method == method].set_index("date").loc[dates]
        for col in fields:
            if col not in rows:
                continue
            values = bundle[col][index]
            expected = values[:, 0] if col == "energy_start_kwh" else values[:, -1] if col == "energy_end_kwh" else values.sum(axis=1)
            close(f"{method}_archive_daily_{col}", rows[col], expected)
        if method == "no_storage":
            close("no_storage_charge_and_discharge_zero", c + d, np.zeros_like(c))
            close("no_storage_soc_constant", soc, np.full_like(soc, 6000))
    bundle.close()


def check_workbook(frame: pd.DataFrame) -> None:
    original = openpyxl.load_workbook(RAW / "附件5" / "result2.xlsx", read_only=True, data_only=True)
    wb = openpyxl.load_workbook(BASE / "result2.xlsx", read_only=False, data_only=True)
    record("result2_sheet_names_preserved", wb.sheetnames == original.sheetnames)
    plan, battery, emergency = wb.worksheets
    record("result2_plan_dimensions", (plan.max_row, plan.max_column) == (335, 147),
           actual_rows=plan.max_row, actual_columns=plan.max_column)
    record("result2_plan_interval_labels_corrected", [normalize_interval(plan.cell(1, k + 2).value) for k in range(144)]
           == [interval(k, k + 1) for k in range(144)])
    record("result2_plan_totals_headers", [plan.cell(1, k).value for k in (146, 147)]
           == [original.worksheets[0].cell(1, k).value for k in (146, 147)])
    by_date = list(frame.groupby("date", sort=False))
    record("result2_plan_dates", [iso_date(plan.cell(k + 2, 1).value) for k in range(334)] == [date for date, _ in by_date])
    plan_values = np.asarray([[plan.cell(row, col).value for col in range(2, 146)] for row in range(2, 336)], dtype=float)
    close("result2_every_plan_cell", plan_values.ravel(), frame.grid_plan_kwh)
    close("result2_daily_plan_energy_sum", [plan.cell(row, 146).value for row in range(2, 336)],
          [day.grid_plan_kwh.sum() for _, day in by_date])
    close("result2_daily_plan_cost_excludes_emergency", [plan.cell(row, 147).value for row in range(2, 336)],
          [(day.grid_plan_kwh * day.price_yuan_per_kwh).sum() for _, day in by_date])
    record("result2_battery_dimensions", (battery.max_row, battery.max_column) == (2005, 6),
           actual_rows=battery.max_row, actual_columns=battery.max_column)
    record("result2_battery_headers_preserved", [battery.cell(1, col).value for col in range(1, 7)]
           == [original.worksheets[1].cell(1, col).value for col in range(1, 7)])
    actual_charge, actual_discharge, expected_charge, expected_discharge = [], [], [], []
    initial, final, expected_initial, expected_final = [], [], [], []
    date_ok = labels_ok = state_labels_ok = True
    expected_emergency: list[tuple[str, str, float]] = []
    for day_index, (date, day) in enumerate(by_date):
        row = 2 + 6 * day_index
        date_ok &= iso_date(battery.cell(row, 1).value) == date
        for k in range(6):
            labels_ok &= normalize_interval(battery.cell(row + k, 2).value) == interval(k * 24, (k + 1) * 24)
            actual_charge.append(battery.cell(row + k, 3).value)
            actual_discharge.append(battery.cell(row + k, 4).value)
            expected_charge.append(day.charge_kwh.iloc[k * 24:(k + 1) * 24].sum())
            expected_discharge.append(day.discharge_kwh.iloc[k * 24:(k + 1) * 24].sum())
        first_label, last_label = battery.cell(row, 5).value, battery.cell(row + 1, 5).value
        state_labels_ok &= first_label in {dt.time(0, 0), "0:00", "00:00", "00:00:00"} and last_label == "24:00"
        initial.append(battery.cell(row, 6).value)
        final.append(battery.cell(row + 1, 6).value)
        expected_initial.append(day.energy_start_kwh.iloc[0])
        expected_final.append(day.energy_end_kwh.iloc[-1])
        q = day.emergency_kwh.to_numpy()
        padded = np.r_[False, q > ACTIVE_TOL, False].astype(int)
        starts = np.flatnonzero(np.diff(padded) == 1)
        ends = np.flatnonzero(np.diff(padded) == -1)
        if not len(starts):
            expected_emergency.append((date, "无", 0.0))
        else:
            expected_emergency.extend((date, interval(int(start), int(end)), float(q[start:end].sum()))
                                      for start, end in zip(starts, ends))
    record("result2_battery_dates", date_ok)
    record("result2_battery_4_hour_interval_labels", labels_ok)
    record("result2_battery_start_end_time_labels", state_labels_ok)
    close("result2_battery_all_4_hour_charge_totals", actual_charge, expected_charge)
    close("result2_battery_all_4_hour_discharge_totals", actual_discharge, expected_discharge)
    close("result2_battery_all_day_initial_soc", initial, expected_initial)
    close("result2_battery_all_day_final_soc", final, expected_final)
    record("result2_emergency_headers_preserved", [emergency.cell(1, col).value for col in range(1, 4)]
           == [original.worksheets[2].cell(1, col).value for col in range(1, 4)])
    actual_emergency = []
    current_date = None
    for row in emergency.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            current_date = iso_date(row[0])
        actual_emergency.append((current_date, normalize_interval(row[1]), row[2]))
    record("result2_emergency_exact_contiguous_interval_coverage",
           [(d, t) for d, t, _ in actual_emergency] == [(d, t) for d, t, _ in expected_emergency],
           actual_intervals=len(actual_emergency), expected_intervals=len(expected_emergency))
    close("result2_emergency_all_interval_energy_totals", [v for _, _, v in actual_emergency],
          [v for _, _, v in expected_emergency])
    close("result2_emergency_grand_total", [sum(float(v) for _, _, v in actual_emergency)],
          [float(frame.emergency_kwh.sum())], tol=2e-3)
    record("result2_no_placeholder_ellipsis", all(str(cell.value).strip() not in {"⋮", "⁝", "…", "..."}
           for ws in wb for row in ws for cell in row))
    wb.close()
    original.close()


def check_causality(frame: pd.DataFrame) -> None:
    """Challenge future access with changed and physically absent observations.

    Forecast caches are rebuilt after perturbation. For the truncation variant,
    current/future load and PV rows do not exist, so accidentally reading them
    cannot silently pass. Calendar dates and the known tariff are retained.
    """
    from forecast import (ForecastBundle, WARMUP_DAYS, YearData, build_forecasts,
                          forecast_at, load_data, make_scenarios, select_methods)
    from optimization import solve_dispatch

    data = load_data()
    base_bundle = build_forecasts(data)
    base_methods, base_scores = select_methods(data)
    for date in ("2025-02-01", "2025-06-01", "2025-12-30"):
        origin = data.dates.index(dt.date.fromisoformat(date))
        horizon = 2
        tag = date.replace("-", "_")
        altered_load, altered_pv = data.load_kw.copy(), data.pv_kw.copy()
        # Distinct transformations avoid preserving the original net load.
        altered_load[origin:] = 19 * altered_load[origin:] + 123456.0
        altered_pv[origin:] = 7 * altered_pv[origin:] + 765432.0
        altered = YearData(data.dates, data.price.copy(), altered_load, altered_pv, data.audit)
        altered_bundle = build_forecasts(altered)
        changed_methods, changed_scores = select_methods(altered)
        record(f"causality_{tag}_january_model_selection_unchanged",
               changed_methods == base_methods and changed_scores == base_scores)

        # Keep the known calendar up to today's origin; realized arrays stop
        # yesterday. This requires no current realized quantity whatsoever.
        truncated = YearData(data.dates[:origin + 1], data.price.copy(),
                             data.load_kw[:origin].copy(), data.pv_kw[:origin].copy(), data.audit)
        load_cache = np.full((origin + 1, 2, 144), np.nan)
        pv_cache = np.full_like(load_cache, np.nan)
        for historical_origin in range(WARMUP_DAYS, origin + 1):
            load_cache[historical_origin], pv_cache[historical_origin] = forecast_at(
                truncated, historical_origin, 2, base_methods)
        truncated_bundle = ForecastBundle(load_cache, pv_cache, base_methods, {})
        scenarios = []
        for variant, variant_data, variant_bundle in [
            ("original", data, base_bundle),
            ("future_perturbed", altered, altered_bundle),
            ("future_rows_removed", truncated, truncated_bundle),
        ]:
            for label in ("load_kwh", "pv_kwh"):
                close(f"causality_{tag}_{variant}_{label}_forecast",
                      getattr(variant_bundle, label)[origin, :horizon],
                      getattr(base_bundle, label)[origin, :horizon], tol=0.0)
            scenario = make_scenarios(variant_data, variant_bundle, origin, horizon)
            scenarios.append(scenario)
            record(f"causality_{tag}_{variant}_residual_targets_strictly_in_past",
                   bool(np.max(scenario.history_indices + horizon - 1) < origin),
                   latest_observed_index=int(np.max(scenario.history_indices + horizon - 1)), origin_index=origin)
            if variant != "original":
                for label in ("load_kwh", "pv_kwh", "weights", "history_indices"):
                    close(f"causality_{tag}_{variant}_scenario_{label}", getattr(scenario, label),
                          getattr(scenarios[0], label), tol=0.0)
        initial = float(frame.loc[frame.date == date, "energy_start_kwh"].iloc[0])
        plans = [solve_dispatch((scenario.load_kwh - scenario.pv_kwh).reshape(len(scenario.weights), -1),
                                scenario.weights, np.tile(data.price, horizon), initial, 6000)
                 for scenario in scenarios]
        for index, variant in [(1, "future_perturbed"), (2, "future_rows_removed")]:
            for key in ("grid", "charge", "discharge", "energy", "net_quantile"):
                close(f"causality_{tag}_{variant}_plan_{key}", plans[index][key], plans[0][key], tol=1e-7)
        day = frame[frame.date == date]
        for key, col in [("grid", "grid_plan_kwh"), ("charge", "charge_kwh"),
                         ("discharge", "discharge_kwh"), ("energy", "energy_end_kwh")]:
            close(f"causality_{tag}_recomputed_plan_matches_exported_{key}", plans[0][key][:144], day[col])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-causality", action="store_true", help="Only audit exported files and physics.")
    args = parser.parse_args()
    frame = pd.read_csv(BASE / "dispatch_detail.csv", dtype={"date": str})
    raw = raw_arrays()
    totals = check_trace(frame, raw)
    check_daily(frame)
    check_all_method_traces(frame, raw)
    check_workbook(frame)
    if not args.skip_causality:
        check_causality(frame)
    report = {
        "all_passed": all(check["passed"] for check in CHECKS),
        "passed_checks": sum(check["passed"] for check in CHECKS),
        "total_checks": len(CHECKS),
        "independent_totals": totals,
        "scope": "All four methods: independently reconstruct every slot from raw actual data and check physical feasibility, costs, and daily totals. Selected method: additionally audit CSV and every workbook output.",
        "limits": "Numerical audit establishes exported schedule feasibility and accounting consistency; it does not prove optimality among every possible adaptive stochastic policy.",
        "causality_requested": not args.skip_causality,
        "checks": CHECKS,
    }
    (BASE / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "checks"}, ensure_ascii=False, indent=2))
    failed = [check for check in CHECKS if not check["passed"]]
    if failed:
        print(json.dumps(failed, ensure_ascii=False, indent=2))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
