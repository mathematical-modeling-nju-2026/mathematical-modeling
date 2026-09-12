"""问题2 结果导出：按附件5/result2.xlsx 模板写出三个工作表。

列对齐说明
    模板"计划购电量"表头首列为 '0:10-0:20'、末列为 '0:00+1-0:10+1'，而附件1/附件2
    的 144 个时段标签是 '0:10' ... '24:00'（区间末）。二者相差一格。本实现按
    **列序 = 附件时段序** 对齐（第 i 列就是附件第 i 个 10 分钟时段），
    并直接复用模板表头文本，保证与模板逐列对应。

紧急购电的合并
    逐段列出 144 个时段会淹没结果，且题目表4 的示例本来就是按连续区间填写的
    （"13:00-13:30"）。故把每天连续的紧急购电时段合并成区间后再填写。
"""

from __future__ import annotations

from datetime import datetime, time

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font

from q2_data import DT, T, template_labels

BLOCKS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
          "12:00-16:00", "16:00-20:00", "20:00-24:00"]

TIMELABEL = [f"{h:02d}:{m:02d}" for h in range(24) for m in (0, 10, 20, 30, 40, 50)]


def _end_label(t):
    """第 t 段（1-based）的区间末标签，t=144 -> '24:00'。"""
    m = t * 10
    return f"{m // 60:02d}:{m % 60:02d}"


def merge_intervals(e, tol=1e-6):
    """把一天内连续的紧急购电时段合并成区间，返回 [(起, 止, 电量), ...]。"""
    out = []
    t = 0
    while t < T:
        if e[t] > tol:
            t0 = t
            q = 0.0
            while t < T and e[t] > tol:
                q += e[t]
                t += 1
            out.append((t0, t, q))
        else:
            t += 1
    return out


def _interval_label(t0, t1):
    return f"{_end_label(t0)}-{_end_label(t1)}"


def export_result2(res, dates, price, path):
    """写出 result2.xlsx（三个工作表：计划购电量 / 充放电量 / 紧急购电量）。"""
    labels = template_labels()
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    # ---------------- 计划购电量 ----------------
    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["日期\\时间"] + labels + ["全天购电量", "全天购电费"])
    for c in ws[1]:
        c.font = bold
        c.alignment = Alignment(horizontal="center")
    for i in range(len(dates)):
        row = [datetime(dates[i].year, dates[i].month, dates[i].day)] + \
              [round(float(v), 4) for v in res.g[i]] + \
              [round(float(res.g[i].sum()), 4),
               round(float(res.plan_cost[i]), 2)]
        ws.append(row)
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "B2"

    # ---------------- 充放电量 ----------------
    ws2 = wb.create_sheet("充放电量")
    ws2.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for c in ws2[1]:
        c.font = bold
    for i in range(len(dates)):
        day = datetime(dates[i].year, dates[i].month, dates[i].day)
        for b, name in enumerate(BLOCKS):
            t0, t1 = b * 36, (b + 1) * 36
            ws2.append([day if b == 0 else None, name,
                        round(float(res.c[i, t0:t1].sum()), 4),
                        round(float(res.d[i, t0:t1].sum()), 4),
                        time(0, 0) if b == 0 else ("24:00" if b == 1 else None),
                        round(float(res.E_start[i]), 4) if b == 0
                        else (round(float(res.E_end[i]), 4) if b == 1 else None)])
    for col, w in zip("ABCDEF", (12, 14, 12, 12, 10, 12)):
        ws2.column_dimensions[col].width = w

    # ---------------- 紧急购电量 ----------------
    ws3 = wb.create_sheet("紧急购电量")
    ws3.append(["日期", "购电时间段", "购电量"])
    for c in ws3[1]:
        c.font = bold
    for i in range(len(dates)):
        ivs = merge_intervals(res.e[i])
        day = datetime(dates[i].year, dates[i].month, dates[i].day)
        if not ivs:
            ws3.append([day, "—", 0.0])
            continue
        for j, (t0, t1, q) in enumerate(ivs):
            ws3.append([day if j == 0 else None, _interval_label(t0, t1),
                        round(float(q), 4)])
    for col, w in zip("ABC", (12, 20, 12)):
        ws3.column_dimensions[col].width = w

    wb.save(path)
    return path


def export_daily_summary(res, dates, path):
    """逐日汇总 CSV，含计划/紧急/合计费用与储能状态，便于论文取数。"""
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["日期", "计划购电量kWh", "紧急购电量kWh", "计划购电费元",
                    "紧急购电费元", "合计购电费元", "紧急时段数",
                    "0:00储电量kWh", "24:00储电量kWh", "充电量kWh", "放电量kWh",
                    "计划净供给kWh", "弃置电量kWh"])
        for i, d in enumerate(dates):
            w.writerow([
                d.strftime("%Y-%m-%d"),
                f"{res.g[i].sum():.4f}", f"{res.e[i].sum():.4f}",
                f"{res.plan_cost[i]:.2f}", f"{res.emg_cost[i]:.2f}",
                f"{res.plan_cost[i] + res.emg_cost[i]:.2f}",
                int(np.sum(res.e[i] > 1e-6)),
                f"{res.E_start[i]:.4f}", f"{res.E_end[i]:.4f}",
                f"{res.c[i].sum():.4f}", f"{res.d[i].sum():.4f}",
                f"{res.z[i].sum():.4f}",
                f"{np.maximum(res.z[i] - res.N[i] * DT, 0.0).sum():.4f}",
            ])
    return path


def export_detail(res, dates, path):
    """逐时段明细 CSV（334 天 × 144 段），供复核。"""
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["日期", "时段", "计划购电量kWh", "充电量kWh", "放电量kWh",
                    "计划净供给kWh", "净负载预测kWh", "紧急购电量kWh", "储电量kWh"])
        for i, d in enumerate(dates):
            for t in range(T):
                w.writerow([d.strftime("%Y-%m-%d"), _end_label(t + 1),
                            f"{res.g[i, t]:.4f}", f"{res.c[i, t]:.4f}",
                            f"{res.d[i, t]:.4f}", f"{res.z[i, t]:.4f}",
                            f"{res.fc[i, t]:.4f}", f"{res.e[i, t]:.4f}",
                            f"{res.E_end[i]:.4f}" if t == T - 1 else ""])
    return path
