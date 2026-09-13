"""Independently verify window experiments against raw attachments and exports.

No optimizer, forecast, or settlement routine is imported here.  Monetary and
physical checks are recalculated from the exported decisions and raw inputs.
"""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, time
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parents[1] / "results"
DT = 1 / 6
T = 144
TOL = 1e-5


def _date(value: Any) -> str:
    return pd.Timestamp(value).strftime("%Y-%m-%d")


def _time_label(slot: int) -> str:
    minute = slot * 10
    return f"{minute // 60:02d}:{minute % 60:02d}"


def _minutes(value: Any) -> int:
    if isinstance(value, (time, datetime)):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    if text in {"0:00+1", "00:00+1"}:
        return 1440
    match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::00)?", text)
    if not match:
        raise ValueError(f"Unrecognized time: {value!r}")
    return int(match[1]) * 60 + int(match[2])


def _interval(value: Any) -> tuple[int, int]:
    parts = re.split(r"[-–—~～]", str(value).strip())
    if len(parts) != 2:
        raise ValueError(f"Unrecognized interval: {value!r}")
    return _minutes(parts[0]), _minutes(parts[1])


def _max_abs(values: Any) -> float:
    arr = np.asarray(values, dtype=float)
    return float(np.max(np.abs(arr))) if arr.size else 0.0


def _read_raw(raw_dir: Path) -> tuple[np.ndarray, dict[str, np.ndarray]]:
    """Read XLSX independently; attachment rows denote interval end times."""
    attachments = raw_dir if (raw_dir / "附件1.xlsx").exists() else raw_dir / "附件"
    wb1 = openpyxl.load_workbook(attachments / "附件1.xlsx", read_only=True, data_only=True)
    rows1 = list(wb1.active.iter_rows(min_row=2, values_only=True))
    wb1.close()
    if len(rows1) != T:
        raise ValueError(f"Attachment 1 has {len(rows1)} data rows instead of 144")
    if any(_minutes(row[0]) % 1440 != ((slot + 1) * 10) % 1440
           for slot, row in enumerate(rows1)):
        raise ValueError("Attachment 1 interval-end labels are not consecutive ten-minute times")
    price = np.asarray([row[1] for row in rows1], dtype=float)
    wb2 = openpyxl.load_workbook(attachments / "附件2.xlsx", read_only=True, data_only=True)
    channels: list[dict[str, np.ndarray]] = []
    for name in ["小区负载", "光伏发电实际功率"]:
        header = next(wb2[name].iter_rows(min_row=1, max_row=1, values_only=True))
        if len(header) != T + 1 or any(_minutes(value) % 1440 != ((slot + 1) * 10) % 1440
                                      for slot, value in enumerate(header[1:])):
            raise ValueError(f"Attachment 2 time columns are inconsistent: {name}")
        data: dict[str, np.ndarray] = {}
        for row in wb2[name].iter_rows(min_row=2, values_only=True):
            day = _date(row[0])
            values = np.asarray(row[1:], dtype=float)
            if len(values) != T or not np.isfinite(values).all():
                raise ValueError(f"Invalid raw row {name}: {day}")
            if day in data:
                raise ValueError(f"Duplicate raw date {name}: {day}")
            data[day] = values
        channels.append(data)
    wb2.close()
    if channels[0].keys() != channels[1].keys():
        raise ValueError("Attachment 2 load/PV dates do not match")
    net = {day: (channels[0][day] - channels[1][day]) * DT for day in channels[0]}
    return price, net


def check_schedule(
    detail: pd.DataFrame,
    daily: pd.DataFrame,
    expected_initial: float,
    raw_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Check all 334 reporting days, accounting, feasibility, and date metadata.

    Historical dates are necessary checks, not a proof that hidden fitting code
    is causal; a separate perturbation test must recompute that pipeline.
    """
    errors: list[str] = []
    metrics: dict[str, Any] = {}

    def check(name: str, condition: Any, metric: Any = None) -> None:
        if metric is not None:
            metrics[name] = metric
        if not bool(condition):
            errors.append(name)

    required_detail = {
        "date", "slot", "start", "end", "price_yuan_per_kwh", "forecast_net_kwh",
        "actual_net_kwh", "g_kwh", "c_kwh", "d_kwh", "energy_before_kwh",
        "energy_after_kwh", "z_kwh", "emergency_kwh", "dump_kwh",
        "planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan",
    }
    required_daily = {
        "date", "planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan",
        "g_kwh", "c_kwh", "d_kwh", "emergency_kwh", "dump_kwh",
        "energy_start_kwh", "energy_end_kwh", "chosen_candidate", "sample_count",
        "effective_sample_count", "max_history_day", "selection_last_observed_day",
    }
    missing_detail = sorted(required_detail - set(detail.columns))
    missing_daily = sorted(required_daily - set(daily.columns))
    if missing_detail or missing_daily:
        return {"pass": False, "errors": ["missing_columns"],
                "missing_detail_columns": missing_detail, "missing_daily_columns": missing_daily}
    detail, daily = detail.copy(), daily.copy()
    try:
        detail["date"] = detail["date"].map(_date)
        daily["date"] = daily["date"].map(_date)
        num_detail = sorted(required_detail - {"date", "start", "end"})
        num_daily = sorted(required_daily - {"date", "chosen_candidate", "max_history_day",
                                            "selection_last_observed_day"})
        for col in num_detail:
            detail[col] = pd.to_numeric(detail[col], errors="raise")
        for col in num_daily:
            daily[col] = pd.to_numeric(daily[col], errors="raise")
        finite = bool(np.isfinite(detail[num_detail].to_numpy()).all()
                      and np.isfinite(daily[num_daily].to_numpy()).all())
        check("finite_numeric_values", finite)
        if not finite or detail.empty or daily.empty:
            return {"pass": False, "errors": errors + (["empty_output"] if detail.empty or daily.empty else []),
                    "metrics": metrics}
    except (TypeError, ValueError) as exc:
        return {"pass": False, "errors": ["invalid_column_values"], "detail": str(exc)}

    expected_dates = pd.date_range("2025-02-01", "2025-12-31", freq="D").strftime("%Y-%m-%d").tolist()
    actual_dates = detail["date"].drop_duplicates().tolist()
    check("report_date_range", actual_dates == expected_dates)
    check("daily_date_range", daily["date"].tolist() == expected_dates)
    check("report_row_count", len(detail) == len(expected_dates) * T, int(len(detail)))
    check("daily_row_count", len(daily) == len(expected_dates), int(len(daily)))
    check("unique_date_slot", not detail.duplicated(["date", "slot"]).any())
    check("unique_daily_date", not daily["date"].duplicated().any())
    check("chronological_rows", list(zip(detail["date"], detail["slot"])) ==
          sorted(zip(detail["date"], detail["slot"])))
    slot_errors = 0
    for _, frame in detail.groupby("date", sort=False):
        slot_errors += int(frame["slot"].tolist() != list(range(T)))
    check("complete_slots_per_day", slot_errors == 0, slot_errors)
    try:
        start = detail["start"].map(_minutes).to_numpy()
        end = detail["end"].map(_minutes).to_numpy()
        slots = detail["slot"].to_numpy()
        check("ten_minute_interval_labels", np.array_equal(start, slots * 10)
              and np.array_equal(end, (slots + 1) * 10))
    except (TypeError, ValueError) as exc:
        check("ten_minute_interval_labels", False, str(exc))

    g, c, d = (detail[col].to_numpy(dtype=float) for col in ["g_kwh", "c_kwh", "d_kwh"])
    eb, ea = (detail[col].to_numpy(dtype=float) for col in ["energy_before_kwh", "energy_after_kwh"])
    p = detail["price_yuan_per_kwh"].to_numpy(dtype=float)
    n = detail["actual_net_kwh"].to_numpy(dtype=float)
    z = g + d - c
    em = np.maximum(n - z, 0)
    dump = np.maximum(z - n, 0)
    expected_values = {
        "z_kwh": z,
        "emergency_kwh": em,
        "dump_kwh": dump,
        "planned_cost_yuan": p * g,
        "emergency_cost_yuan": 5 * p * em,
        "total_cost_yuan": p * g + 5 * p * em,
    }
    for col, expected in expected_values.items():
        deviation = _max_abs(detail[col].to_numpy() - expected)
        check(f"{col}_max_error", deviation <= TOL, deviation)
    residual = _max_abs(ea - eb - 0.9 * c + d / 0.9)
    check("soc_equation_max_error_kwh", residual <= TOL, residual)
    residual = _max_abs(eb[1:] - ea[:-1])
    check("soc_continuity_max_error_kwh", residual <= TOL, residual)
    check("initial_soc_kwh", abs(eb[0] - expected_initial) <= TOL, float(eb[0]))
    check("final_soc_kwh", abs(ea[-1] - 6000) <= TOL, float(ea[-1]))
    check("soc_min_kwh", min(eb.min(), ea.min()) >= 1200 - TOL, float(min(eb.min(), ea.min())))
    check("soc_max_kwh", max(eb.max(), ea.max()) <= 10800 + TOL, float(max(eb.max(), ea.max())))
    check("charge_max_kw", c.max() / DT <= 5000 + TOL, float(c.max() / DT))
    check("discharge_max_kw", d.max() / DT <= 5000 + TOL, float(d.max() / DT))
    check("nonnegative_grid_charge_discharge", min(g.min(), c.min(), d.min()) >= -TOL)
    check("nonnegative_emergency_dump", min(detail["emergency_kwh"].min(), detail["dump_kwh"].min()) >= -TOL)
    simultaneous = int(((c > 1e-6) & (d > 1e-6)).sum())
    check("simultaneous_charge_discharge_slots", simultaneous == 0, simultaneous)
    margin = float((z - detail["forecast_net_kwh"].to_numpy()).min())
    check("plan_forecast_min_margin_kwh", margin >= -TOL, margin)
    balance = _max_abs(z + detail["emergency_kwh"] - n - detail["dump_kwh"])
    check("actual_balance_max_error_kwh", balance <= TOL, balance)

    try:
        raw_path = Path(raw_dir) if raw_dir is not None else HERE.parents[1] / "data"
        raw_price, raw_net = _read_raw(raw_path)
        reference_price = np.asarray([raw_price[int(s)] for s in detail["slot"]])
        reference_net = np.asarray([raw_net[day][int(slot)] for day, slot in
                                    zip(detail["date"], detail["slot"])])
        deviation = _max_abs(p - reference_price)
        check("raw_price_max_error_yuan_per_kwh", deviation <= 1e-10, deviation)
        deviation = _max_abs(n - reference_net)
        check("raw_net_energy_max_error_kwh", deviation <= TOL, deviation)
    except (KeyError, IndexError, ValueError, OSError) as exc:
        check("raw_attachment_match", False, str(exc))

    aggregate_cols = ["planned_cost_yuan", "emergency_cost_yuan", "total_cost_yuan", "g_kwh",
                      "c_kwh", "d_kwh", "emergency_kwh", "dump_kwh"]
    if not daily["date"].duplicated().any():
        aggregate = detail.groupby("date")[aggregate_cols].sum()
        target = daily.set_index("date")
        common_dates = aggregate.index.intersection(target.index)
        check("detail_daily_matching_dates", len(common_dates) == len(aggregate) == len(target))
        for col in aggregate_cols:
            deviation = _max_abs(aggregate.loc[common_dates, col] - target.loc[common_dates, col])
            check(f"daily_{col}_max_error", deviation <= 1e-4, deviation)
        for expected_col, field, endpoint in [("energy_start_kwh", "energy_before_kwh", "first"),
                                               ("energy_end_kwh", "energy_after_kwh", "last")]:
            measured = detail.groupby("date")[field].agg(endpoint)
            deviation = _max_abs(measured.loc[common_dates] - target.loc[common_dates, expected_col])
            check(f"daily_{expected_col}_max_error", deviation <= TOL, deviation)

    history_violations = []
    selection_violations = []
    for row in daily.itertuples(index=False):
        for attr, failures, required in [("max_history_day", history_violations, True),
                                          ("selection_last_observed_day", selection_violations, False)]:
            value = getattr(row, attr)
            if pd.isna(value) or str(value).strip() == "":
                if required:
                    failures.append(row.date)
                continue
            try:
                if _date(value) >= row.date:
                    failures.append(row.date)
            except (TypeError, ValueError):
                failures.append(row.date)
    check("history_dates_before_decision", not history_violations, history_violations)
    check("selection_dates_before_decision", not selection_violations, selection_violations)
    counts = daily["sample_count"].to_numpy()
    effective = daily["effective_sample_count"].to_numpy()
    check("valid_sample_counts", bool(np.all(counts >= 1) and np.all(counts <= 56)
                                       and np.all(counts == counts.astype(int))))
    check("effective_sample_count_bounds", bool(np.all(effective >= 1 - TOL)
                                                and np.all(effective <= counts + TOL)))
    check("candidate_names_present", daily["chosen_candidate"].notna().all()
          and daily["chosen_candidate"].astype(str).str.len().gt(0).all())
    metrics["recomputed_total_cost_yuan"] = float((p * g + 5 * p * em).sum())
    metrics["recomputed_emergency_kwh"] = float(em.sum())
    metrics["recomputed_dump_kwh"] = float(dump.sum())
    return {"pass": not errors, "errors": errors, "metrics": metrics,
            "scope": "Exported physical and accounting constraints; independent raw data; historical date metadata. "
                     "Causality of fitting and selection requires the separate future-perturbation audit."}


def check_workbook(path: str | Path, detail: pd.DataFrame) -> dict[str, Any]:
    """Read all three result2 worksheets back, including interval alignment."""
    errors: list[str] = []
    deviations: dict[str, float] = {}

    def numeric(name: str, actual: Any, expected: float, tolerance: float = 5.1e-5) -> None:
        try:
            delta = abs(float(actual) - float(expected))
            if not np.isfinite(delta):
                raise ValueError("nonfinite cell")
        except (TypeError, ValueError):
            delta = float("inf")
        if np.isfinite(delta):
            deviations[name] = max(deviations.get(name, 0), delta)
        if delta > tolerance and name not in errors:
            errors.append(name)

    frame = detail.copy()
    frame["date"] = frame["date"].map(_date)
    frame = frame.sort_values(["date", "slot"])
    days = list(frame.groupby("date", sort=True))
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        required = ["计划购电量", "充放电量", "紧急购电量"]
        if any(name not in wb.sheetnames for name in required):
            wb.close()
            return {"pass": False, "errors": ["missing_worksheets"]}

        plan = list(wb[required[0]].iter_rows(values_only=True))
        if len(plan) != len(days) + 1 or len(plan[0]) != T + 3:
            errors.append("plan_table_dimensions")
        expected_labels = [f"{_time_label(t)}-{_time_label(t + 1)}" for t in range(T)]
        if list(plan[0][1:T + 1]) != expected_labels:
            errors.append("plan_interval_headers")
        for row, (day, group) in zip(plan[1:], days):
            if _date(row[0]) != day:
                errors.append(f"plan_date:{day}")
            for value, expected in zip(row[1:T + 1], group["g_kwh"]):
                numeric("plan_slot_energy_max_error", value, expected)
            numeric("plan_daily_energy_max_error", row[T + 1], group["g_kwh"].sum())
            numeric("plan_daily_cost_max_error", row[T + 2],
                    (group["g_kwh"] * group["price_yuan_per_kwh"]).sum(), 0.0051)

        storage = list(wb[required[1]].iter_rows(values_only=True))
        if len(storage) != 1 + 6 * len(days):
            errors.append("storage_table_dimensions")
        for i, (day, group) in enumerate(days):
            for block in range(6):
                ix = 1 + 6 * i + block
                if ix >= len(storage):
                    continue
                row = storage[ix]
                if block == 0 and _date(row[0]) != day:
                    errors.append(f"storage_date:{day}")
                if _interval(row[1]) != (block * 240, (block + 1) * 240):
                    errors.append(f"storage_interval:{day}:{block}")
                slots = group.iloc[block * 24:(block + 1) * 24]
                numeric("storage_charge_block_max_error", row[2], slots["c_kwh"].sum())
                numeric("storage_discharge_block_max_error", row[3], slots["d_kwh"].sum())
                if block < 2:
                    expected_time = 0 if block == 0 else 1440
                    if _minutes(row[4]) != expected_time:
                        errors.append(f"storage_soc_time:{day}:{block}")
                    expected_soc = group["energy_before_kwh"].iloc[0] if block == 0 else group["energy_after_kwh"].iloc[-1]
                    numeric("storage_soc_max_error", row[5], expected_soc)

        expected_emergency: list[tuple[str, tuple[int, int] | None, float]] = []
        for day, group in days:
            emergency = group["emergency_kwh"].to_numpy()
            cursor = 0
            found = False
            while cursor < T:
                if emergency[cursor] <= 1e-6:
                    cursor += 1
                    continue
                start, amount = cursor, 0.0
                while cursor < T and emergency[cursor] > 1e-6:
                    amount += emergency[cursor]
                    cursor += 1
                expected_emergency.append((day, (start * 10, cursor * 10), amount))
                found = True
            if not found:
                expected_emergency.append((day, None, 0.0))
        emergency_rows = list(wb[required[2]].iter_rows(min_row=2, values_only=True))
        if len(emergency_rows) != len(expected_emergency):
            errors.append("emergency_interval_row_count")
        active_date = None
        for row, (day, expected_interval, amount) in zip(emergency_rows, expected_emergency):
            if row[0] is not None:
                active_date = _date(row[0])
            if active_date != day:
                errors.append(f"emergency_date:{day}")
            if expected_interval is None:
                if row[1] not in {None, "", "—", "-", "无"}:
                    errors.append(f"emergency_empty_label:{day}")
            elif _interval(row[1]) != expected_interval:
                errors.append(f"emergency_interval:{day}:{expected_interval}")
            numeric("emergency_interval_energy_max_error", row[2], amount)
        wb.close()
    except (OSError, IndexError, TypeError, ValueError) as exc:
        errors.append(f"workbook_read_error:{exc}")
    return {"pass": not errors, "errors": errors, "max_deviations": deviations,
            "workbook": str(path)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=HERE)
    parser.add_argument("--raw-dir", type=Path, default=HERE.parents[1] / "data")
    parser.add_argument("--expected-initial", type=float, default=8550.0)
    args = parser.parse_args()
    reports: dict[str, Any] = {}
    for variant in sorted((args.root / "variants").glob("*")):
        detail_path = variant / "schedule_detail.csv.gz"
        daily_path = variant / "daily_summary.csv"
        if not detail_path.exists() or not daily_path.exists():
            continue
        detail = pd.read_csv(detail_path)
        daily = pd.read_csv(daily_path)
        report = {"schedule": check_schedule(detail, daily, args.expected_initial, args.raw_dir)}
        workbook = variant / "result2.xlsx"
        if workbook.exists():
            report["workbook"] = check_workbook(workbook, detail)
        report["pass"] = all(item["pass"] for item in report.values())
        reports[variant.name] = report
        print(f"{variant.name}: {'PASS' if report['pass'] else 'FAIL'}", flush=True)
    result = {"pass": bool(reports) and all(item["pass"] for item in reports.values()),
              "expected_initial_kwh": args.expected_initial, "variants": reports}
    (args.root / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Verified {len(reports)} variants. Overall: {'PASS' if result['pass'] else 'FAIL'}")
    if not result["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
