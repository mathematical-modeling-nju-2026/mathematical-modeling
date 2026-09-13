"""Independent raw-XLSX, physical and cash-accounting audit of Q4-2 exports.

No optimizer, forecast or settlement function is used for these checks.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parents[1] / "results"
REPO = HERE.parents[2]
RAW = REPO/'data'/'附件'
sys.path.insert(0, str(REPO/'Q2'/'code'))
from verify_window import check_workbook  # independent Excel parser, no model imports


def raw_arrays():
    book = openpyxl.load_workbook(RAW/'附件2.xlsx', read_only=True, data_only=True)
    load_rows = list(book['小区负载'].iter_rows(min_row=2, values_only=True))
    pv_rows = list(book['光伏发电实际功率'].iter_rows(min_row=2, values_only=True))
    book.close()
    book = openpyxl.load_workbook(RAW/'附件4.xlsx', read_only=True, data_only=True)
    price_rows = list(book.active.iter_rows(min_row=2, values_only=True))
    book.close()
    expected_dates = list(pd.date_range('2025-01-01', '2025-12-31'))
    for rows in (load_rows, pv_rows, price_rows):
        if [pd.Timestamp(r[0]) for r in rows] != expected_dates:
            raise ValueError('Raw dates do not cover the year exactly')
    load, pv, price = [np.asarray([r[1:145] for r in rows], float)
                       for rows in (load_rows, pv_rows, price_rows)]
    if any(a.shape != (365, 144) or not np.isfinite(a).all() for a in (load, pv, price)):
        raise ValueError('Invalid raw data')
    return (load-pv)/6, price


def check_schedule(frame, daily, net, price, start_day=31, initial=8550., terminal=6000.):
    errors, metrics = [], {}

    def check(name, passed, value=None):
        if value is not None:
            metrics[name] = value
        if not bool(passed):
            errors.append(name)

    def near(name, x, y, tol=1e-5):
        difference = float(np.max(np.abs(np.asarray(x)-np.asarray(y))))
        check(name, np.isfinite(difference) and difference <= tol, difference)

    end_day = 31 if start_day == 0 else 365
    days = end_day-start_day
    expected_dates = pd.date_range('2025-01-01', periods=365)[start_day:end_day].strftime('%Y-%m-%d')
    check('row_count', len(frame) == days*144, int(len(frame)))
    if len(frame) != days*144:
        return dict(pass_check=False, errors=errors, metrics=metrics)
    check('date_order', frame.date.tolist() == np.repeat(expected_dates, 144).tolist())
    check('slot_order', frame.slot.tolist() == np.tile(np.arange(144), days).tolist())
    labels = [f'{t//6:02d}:{t%6*10:02d}' for t in range(145)]
    check('interval_start', frame.start.tolist() == labels[:144]*days)
    check('interval_end', frame.end.tolist() == labels[1:]*days)
    check('daily_date_order', daily.date.tolist() == expected_dates.tolist())
    g, c, d, eb, ea, p, n = (frame[k].to_numpy(float) for k in
        ('g_kwh', 'c_kwh', 'd_kwh', 'energy_before_kwh', 'energy_after_kwh',
         'price_yuan_per_kwh', 'actual_net_kwh'))
    check('finite_values', all(np.isfinite(a).all() for a in (g,c,d,eb,ea,p,n)))
    near('raw_price_max_error', p, price[start_day:end_day].ravel(), 1e-12)
    near('raw_net_load_max_error', n, net[start_day:end_day].ravel())
    near('soc_balance_max_error', ea, eb+.9*c-d/.9)
    near('soc_continuity_max_error', eb[1:], ea[:-1])
    near('initial_soc_error', eb[0], initial)
    near('terminal_soc_error', ea[-1], terminal)
    check('soc_lower_bound', min(eb.min(), ea.min()) >= 1200-1e-5, float(min(eb.min(), ea.min())))
    check('soc_upper_bound', max(eb.max(), ea.max()) <= 10800+1e-5, float(max(eb.max(), ea.max())))
    check('nonnegative_flows', min(g.min(), c.min(), d.min()) >= -1e-6)
    check('charge_power', c.max()*6 <= 5000+1e-5, float(c.max()*6))
    check('discharge_power', d.max()*6 <= 5000+1e-5, float(d.max()*6))
    overlap = int(((c > 1e-6) & (d > 1e-6)).sum())
    check('simultaneous_charge_discharge', overlap == 0, overlap)
    z = g+d-c
    emergency = np.maximum(n-z, 0.)
    dump = np.maximum(z-n, 0.)
    near('net_supply_error', frame.z_kwh, z)
    near('emergency_error', frame.emergency_kwh, emergency)
    near('surplus_error', frame.dump_kwh, dump)
    near('actual_balance_error', z+frame.emergency_kwh, n+frame.dump_kwh)
    check('point_forecast_floor', np.min(z-frame.forecast_net_kwh) >= -1e-5,
          float(np.min(z-frame.forecast_net_kwh)))
    near('planned_cost_error', frame.planned_cost_yuan, p*g)
    near('emergency_cost_error', frame.emergency_cost_yuan, 5*p*emergency)
    near('total_cost_error', frame.total_cost_yuan, p*(g+5*emergency))
    aggregate = frame.groupby('date', sort=False)
    for col in ('planned_cost_yuan', 'emergency_cost_yuan', 'total_cost_yuan',
                'g_kwh', 'c_kwh', 'd_kwh', 'emergency_kwh', 'dump_kwh'):
        near('daily_'+col, daily[col], aggregate[col].sum().to_numpy(), 1e-4)
    near('daily_initial_soc', daily.energy_start_kwh, aggregate.energy_before_kwh.first().to_numpy())
    near('daily_final_soc', daily.energy_end_kwh, aggregate.energy_after_kwh.last().to_numpy())
    if 'max_history_day' in daily:
        check('history_strictly_past', bool((daily.max_history_day < daily.date).all()))
        expected_count = np.minimum(56, np.arange(start_day, end_day)-7)
        near('scenario_count', daily.sample_count, expected_count, 0.)
        check('all_solvers_optimal', bool((daily.solver_status == 0).all()))
    metrics['recomputed_total_cost_yuan'] = float(np.sum(p*(g+5*emergency)))
    metrics['recomputed_emergency_kwh'] = float(emergency.sum())
    return dict(pass_check=not errors, errors=errors, metrics=metrics)


def main():
    net, price = raw_arrays()
    metadata = json.loads((HERE/'run_metadata.json').read_text(encoding='utf-8'))
    initial = metadata['initial_energy_kwh']
    january = pd.read_csv(HERE/'january_warmup.csv.gz')
    january_days = pd.read_csv(HERE/'january_summary.csv')
    reports = dict(january=check_schedule(january, january_days, net, price, 0, 6000., initial), variants={})
    for folder in sorted((HERE/'variants').iterdir()):
        if not folder.is_dir():
            continue
        detail = pd.read_csv(folder/'schedule_detail.csv.gz')
        daily = pd.read_csv(folder/'daily_summary.csv')
        report = dict(schedule=check_schedule(detail, daily, net, price, initial=initial),
                      workbook=check_workbook(folder/'result4-2.xlsx', detail))
        summary = json.loads((folder/'summary.json').read_text(encoding='utf-8'))
        report['summary_total_error'] = abs(summary['total_cost_yuan']-report['schedule']['metrics']['recomputed_total_cost_yuan'])
        report['pass_check'] = (report['schedule']['pass_check'] and report['workbook']['pass']
                                and report['summary_total_error'] <= 1e-4)
        reports['variants'][folder.name] = report
        print(folder.name, 'PASS' if report['pass_check'] else 'FAIL', flush=True)
    selected = HERE/'variants'/metadata['primary_variant']
    detail = pd.read_csv(selected/'schedule_detail.csv.gz')
    daily = pd.read_csv(selected/'daily_summary.csv')
    injections = {}
    for column, amount in [('price_yuan_per_kwh', .1), ('energy_after_kwh', 10.), ('emergency_kwh', 10.)]:
        broken = detail.copy()
        broken.loc[0, column] += amount
        injections[column] = not check_schedule(broken, daily, net, price, initial=initial)['pass_check']
    reports['injected_errors_rejected'] = injections
    fingerprint_path = HERE/'source_fingerprints.migrated.json'
    if not fingerprint_path.exists():
        fingerprint_path = HERE/'source_fingerprints.json'
    fingerprints = json.loads(fingerprint_path.read_text(encoding='utf-8'))
    reports['source_files_unchanged'] = all(hashlib.sha256((REPO/p).read_bytes()).hexdigest() == h for p,h in fingerprints.items())
    reports['pass_check'] = (reports['january']['pass_check'] and bool(reports['variants'])
        and all(r['pass_check'] for r in reports['variants'].values())
        and all(injections.values()) and reports['source_files_unchanged'])
    (HERE/'verification.json').write_text(json.dumps(reports, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Overall', reports['pass_check'])
    if not reports['pass_check']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
