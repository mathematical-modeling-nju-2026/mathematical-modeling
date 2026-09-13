"""Independent physical, accounting, workbook and information-boundary checks."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from q3_data import HERE, T, load_data


def check_schedule(detail, daily, data, require_year_end=True):
    x=detail
    p=np.tile(data['price'],len(daily))
    g,q,c,d=(x[k].to_numpy() for k in ('original_kwh','adjusted_kwh','charge_kwh','discharge_kwh'))
    day_indices=(pd.to_datetime(x.date)-pd.Timestamp('2025-01-01')).dt.days.to_numpy()
    slots=x.slot.to_numpy(int)
    net=(data['load'][day_indices,slots]-data['pv'][day_indices,slots])/6
    emergency=np.maximum(net+c-d-q,0)
    down=np.maximum(g-q,0); up=np.maximum(q-g,0)
    # Independently use paid delivered base + breach + new purchases.
    cost=p*np.minimum(g,q)+.5*p*down+1.5*p*up+5*p*emergency
    before=x.energy_before_kwh.to_numpy(); after=x.energy_after_kwh.to_numpy()
    balance=q+d+x.emergency_kwh.to_numpy()-net-c-x.dump_kwh.to_numpy()
    sumcols=['planned_cost_yuan','increase_cost_yuan','cancellation_refund_yuan','breach_cost_yuan',
             'adjustment_net_yuan','emergency_cost_yuan','total_cost_yuan','original_kwh','adjusted_kwh',
             'increase_kwh','decrease_kwh','emergency_kwh','dump_kwh','charge_kwh','discharge_kwh']
    summed=x.groupby('date',sort=False)[sumcols].sum().to_numpy()
    errs=dict(
        periods=int(len(x)), duplicate_date_slots=int(x.duplicated(['date','slot']).sum()),
        energy_min_kwh=float(min(after.min(),before.min())), energy_max_kwh=float(max(after.max(),before.max())),
        energy_recursion_max_error=float(np.max(abs(after-before-.9*c+d/.9))),
        chronological_continuity_max_error=float(np.max(abs(after[:-1]-before[1:]))),
        initial_energy_kwh=float(before[0]), final_energy_kwh=float(after[-1]),
        max_charge_kwh=float(c.max()), max_discharge_kwh=float(d.max()),
        min_flow_kwh=float(min(c.min(),d.min(),q.min(),g.min())),
        simultaneous_periods=int(((c>1e-7)&(d>1e-7)).sum()),
        energy_balance_max_error=float(abs(balance).max()),
        emergency_max_error=float(np.max(abs(emergency-x.emergency_kwh.to_numpy()))),
        settlement_per_slot_max_error=float(np.max(abs(cost-x.total_cost_yuan.to_numpy()))),
        daily_sum_max_error=float(np.max(abs(summed-daily[sumcols].to_numpy()))),
        cost_recomputed_yuan=float(cost.sum()),
        altered_past_periods=int(np.sum(x.issue_hour.to_numpy()>slots/6)),
        midnight_block_plan_max_error=float(np.max(abs((q-g)[x.issue_hour.to_numpy()==0]))),
        lp_eq_max=float(daily.lp_eq_max.max()),lp_ineq_max=float(daily.lp_ineq_max.max()),
        alternative_no_refund_cost_same_schedule_yuan=float((cost+p*down).sum()))
    errs['pass']=bool(
        len(x)==len(daily)*T and errs['duplicate_date_slots']==0 and
        errs['energy_min_kwh']>=1200-1e-6 and errs['energy_max_kwh']<=10800+1e-6 and
        errs['energy_recursion_max_error']<1e-6 and errs['chronological_continuity_max_error']<1e-6 and
        abs(before[0]-6000)<1e-6 and (not require_year_end or abs(after[-1]-6000)<1e-6) and
        max(c.max(),d.max())<=5000/6+1e-6 and errs['min_flow_kwh']>=-1e-6 and
        errs['simultaneous_periods']==0 and errs['energy_balance_max_error']<1e-6 and
        errs['emergency_max_error']<1e-6 and errs['settlement_per_slot_max_error']<1e-6 and
        errs['daily_sum_max_error']<1e-5 and errs['altered_past_periods']==0 and
        errs['midnight_block_plan_max_error']<1e-6 and errs['lp_eq_max']<1e-6 and errs['lp_ineq_max']<1e-6)
    return errs


def information_checks(data):
    from q3_forecast import Forecaster
    from q3_model import solve, make_tree
    base=Forecaster(data)
    tests=[]
    for day in (31,171,354):
        original=None; energy=6000.
        for stage in range(4):
            a=solve(base,day,stage,(0,1,2,3),energy,6000.,original)
            changed={k:(v.copy() if isinstance(v,np.ndarray) else v) for k,v in data.items()}
            # Includes all still-unrealized current-day values, all later days,
            # and all forecasts released strictly after the current issue.
            cut=day*T+stage*36
            changed['load'].ravel()[cut:]*=1.7
            changed['pv'].ravel()[cut:]+=321.
            changed['external'][day,stage+1:]+=777.
            changed['external'][day+1:]+=555.
            other=Forecaster(changed)
            b=solve(other,day,stage,(0,1,2,3),energy,6000.,original)
            prediction_error=float(np.max(abs(base.point(day,stage)[stage*36:]-other.point(day,stage)[stage*36:])))
            control_error=max(float(np.max(abs(a[k]-b[k]))) for k in ('original','q','c','d','E'))
            tests.append(dict(day=day,issue_hour=stage*6,prediction_max_change=prediction_error,
                              decision_max_change=control_error,pass_check=control_error<1e-7 and prediction_error<1e-7))
            original=a['original']; energy=a['E'][-1]
    # Test zero-update identity: same optimization returns no adjustment.
    r=solve(base,171,0,(0,),6000.,6000.)
    identity=float(np.max(abs(r['q']-r['original'])))
    result=dict(future_perturbation_tests=tests, no_update_identity_max_error=identity,
                pass_check=all(t['pass_check'] for t in tests) and identity<1e-7)
    return result


def _workbook_slot_order(headers):
    """Map result-table headers to internal slots, including template wraparound."""
    slots=[]
    for label in headers:
        try:
            start=str(label).split('-',1)[0].replace('+1','')
            hour,minute=map(int,start.split(':'))
            slots.append((hour*6+minute//10)%T)
        except (AttributeError,TypeError,ValueError):
            return []
    return slots if len(slots)==T and sorted(slots)==list(range(T)) else []


def workbook_checks(path,detail):
    import openpyxl
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    errors=[]
    for name,col in [('计划购电量','original_kwh'),('调整购电量','adjusted_kwh')]:
        s=w[name]
        slot_order=_workbook_slot_order([s.cell(1,col).value for col in range(2,146)])
        if not slot_order:
            w.close()
            return dict(max_roundtrip_error=float('inf'),pass_check=False,
                        errors=[f'{name}: invalid interval headers'])
        actual=np.array([r[1:145] for r in s.iter_rows(min_row=2,values_only=True)],float)
        expected=detail[col].to_numpy().reshape(-1,T)[:,slot_order]
        errors.append(float(np.max(abs(actual-expected))))
    rows=list(w['充放电量'].iter_rows(min_row=2,values_only=True))
    charge=np.array([r[2] for r in rows]); discharge=np.array([r[3] for r in rows])
    errors.extend([float(abs(charge-detail.charge_kwh.to_numpy().reshape(-1,24).sum(1)).max()),
                   float(abs(discharge-detail.discharge_kwh.to_numpy().reshape(-1,24).sum(1)).max())])
    total_emergency=sum(float(r[2] or 0) for r in w['紧急购电量'].iter_rows(min_row=2,values_only=True))
    errors.append(abs(total_emergency-detail.emergency_kwh.sum()))
    w.close()
    # Workbook values are exported to six decimals; summing rounded blocks can
    # accumulate a few 1e-6 units, so use a rounding-aware tolerance.
    return dict(max_roundtrip_error=float(max(errors)),pass_check=bool(max(errors)<5.1e-5))


def main():
    data=load_data()
    d=pd.read_csv(HERE/'schedule_detail.csv'); days=pd.read_csv(HERE/'daily_summary.csv')
    result=check_schedule(d,days,data)
    result['information_boundary']=information_checks(data)
    result['workbook']=workbook_checks(HERE/'result3.xlsx',d)
    result['pass']=bool(result['pass'] and result['information_boundary']['pass_check'] and result['workbook']['pass_check'])
    alternative=HERE/'comparisons'/'at_0_6_12'
    if (alternative/'result3.xlsx').exists():
        result['alternative_workbook']=workbook_checks(alternative/'result3.xlsx',pd.read_csv(alternative/'schedule_detail.csv'))
        result['pass']=bool(result['pass'] and result['alternative_workbook']['pass_check'])
    (HERE/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
    if not result['pass']: raise SystemExit(1)


if __name__=='__main__': main()
